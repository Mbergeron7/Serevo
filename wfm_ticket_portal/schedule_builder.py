"""
schedule_builder.py
--------------------
WFM TL Rotation Scheduler — Flask-integrated version of TL_Rotation_Scheduler.ipynb.

Implements:
  - SS weekday closing rotation with skip-week logic
  - SS Saturday (primary TL + fill fallback)
  - SS Sunday LOD rotation
  - PS Mon–Thu static closer, Friday rotation, weekend LOD
  - Stat holiday assignment (SS closes / PS opens)
  - PTO substitution with fill notes
  - Text + HTML reports
  - Excel export (6-tab workbook: SS Rotation, PS Rotation, Stat Holidays, PTO, Accommodations, Summary)
"""

import logging
import datetime
from collections import defaultdict
from io import BytesIO

log = logging.getLogger(__name__)

DAYS_ORDER = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


# =========================================================
# MAIN ENTRY POINT
# =========================================================

def build_schedule(
    ss_list,
    ps_list,
    zero_tls=None,
    accommodations=None,
    pto=None,
    holidays=None,
    # Rotation config
    ss_closing_rotation=None,
    ss_sunday_rotation=None,
    ps_friday_rotation=None,
    ps_weekend_tl="",
    ps_monday_thursday_tl="",
    ss_saturday_primary_tl="",
    # Start indexes
    ss_skip_start=0,
    ps_fri_start=0,
    ss_sun_start=0,
    # Schedule date range
    schedule_start=None,
    schedule_end=None,
    # Shift times
    shift_times=None,
):
    """
    Main entry called by Flask. Returns a dict with:
        schedule_ss, schedule_ps  — per-week dicts
        text_report               — plain-text summary
        html_report               — Bootstrap HTML report
        excel_bytes               — BytesIO of the .xlsx workbook
    """
    zero_tls       = zero_tls       or []
    accommodations = accommodations or {}
    pto            = pto            or []
    holidays       = holidays       or []
    shift_times    = shift_times    or _default_shift_times()

    all_tls = {n: "SS" for n in ss_list}
    all_tls.update({n: "PS" for n in ps_list})

    # Defaults for rotations if not provided
    ss_closing_rotation   = ss_closing_rotation   or [t for t in ss_list  if t not in zero_tls]
    ss_sunday_rotation    = ss_sunday_rotation    or ss_closing_rotation
    ps_friday_rotation    = ps_friday_rotation    or [t for t in ps_list  if t not in zero_tls]
    ps_weekend_tl         = ps_weekend_tl         or (ps_list[-1] if ps_list else "")
    ps_monday_thursday_tl = ps_monday_thursday_tl or (ps_list[0]  if ps_list else "")
    ss_saturday_primary   = ss_saturday_primary_tl or (ss_closing_rotation[0] if ss_closing_rotation else "")

    # Date range — default to next Monday + 26 weeks
    if schedule_start:
        try:
            start_dt = datetime.datetime.strptime(schedule_start, "%Y-%m-%d").date()
        except Exception:
            start_dt = _next_monday()
    else:
        start_dt = _next_monday()

    if schedule_end:
        try:
            end_dt = datetime.datetime.strptime(schedule_end, "%Y-%m-%d").date()
        except Exception:
            end_dt = start_dt + datetime.timedelta(weeks=26)
    else:
        end_dt = start_dt + datetime.timedelta(weeks=26)

    # Normalize
    norm_pto      = _normalize_pto(pto)
    norm_holidays = _normalize_holidays(holidays)
    norm_accom    = _normalize_accommodations(accommodations, all_tls)

    mondays = _week_mondays(start_dt, end_dt)

    # Run engines
    ss_weekday = _gen_ss_weekday(mondays, ss_closing_rotation, norm_pto, ss_skip_start)
    ss_sundays = _gen_ss_sunday(mondays, ss_sunday_rotation, norm_pto, ss_sun_start)
    ss_sats    = _gen_ss_saturday(mondays, ss_saturday_primary, ss_closing_rotation, norm_pto)
    ps_sched   = _gen_ps(mondays, ps_friday_rotation, ps_weekend_tl,
                         ps_monday_thursday_tl, norm_pto, ps_fri_start)
    stat_recs  = _assign_stats(norm_holidays, ss_weekday, ps_sched,
                               ss_sats, ss_sundays, shift_times)

    # Build reports
    text_report  = _build_text_report(ss_weekday, ps_sched, ss_sats, ss_sundays,
                                      stat_recs, start_dt, end_dt)
    html_report  = _build_html_report(ss_weekday, ps_sched, ss_sats, ss_sundays,
                                      stat_recs, start_dt, end_dt, shift_times)
    excel_bytes  = _build_excel(ss_weekday, ps_sched, ss_sats, ss_sundays,
                                stat_recs, norm_pto, norm_accom, norm_holidays,
                                all_tls, shift_times)

    return {
        "schedule_ss":  ss_weekday,
        "schedule_ps":  ps_sched,
        "text_report":  text_report,
        "html_report":  html_report,
        "excel_bytes":  excel_bytes,
        # kept for backwards compat
        "summary_ss":       {},
        "summary_ps":       {},
        "rag_fairness_ss":  {},
        "rag_fairness_ps":  {},
    }


# =========================================================
# DATE HELPERS
# =========================================================

def _next_monday():
    today = datetime.date.today()
    days  = (7 - today.weekday()) % 7 or 7
    return today + datetime.timedelta(days=days)

def _week_mondays(start, end):
    # Snap start back to its Monday
    start = start - datetime.timedelta(days=start.weekday())
    result, d = [], start
    while d <= end:
        result.append(d)
        d += datetime.timedelta(weeks=1)
    return result

def _day_name(d):
    return DAYS_ORDER[d.weekday()]


# =========================================================
# PTO / HOLIDAY NORMALIZATION
# =========================================================

def _normalize_pto(pto_list):
    """Return {name: set(dates)}"""
    out = defaultdict(set)
    for entry in pto_list:
        try:
            name  = entry["name"].strip()
            start = datetime.datetime.strptime(entry["start"], "%Y-%m-%d").date()
            end   = datetime.datetime.strptime(entry["end"],   "%Y-%m-%d").date()
            if start > end:
                start, end = end, start
            d = start
            while d <= end:
                out[name].add(d)
                d += datetime.timedelta(days=1)
        except Exception:
            continue
    return dict(out)

def _normalize_holidays(holidays):
    out = []
    for h in holidays:
        try:
            out.append({
                "date": datetime.datetime.strptime(h["date"], "%Y-%m-%d").date(),
                "name": h.get("name", ""),
            })
        except Exception:
            continue
    return out

def _normalize_accommodations(accommodations, all_tls):
    """Normalize free-form accommodation dict from the form."""
    out = {}
    for name, data in accommodations.items():
        if isinstance(data, dict):
            row = {"group": all_tls.get(name, ""), "notes": data.get("notes", "")}
            for day in DAYS_ORDER:
                row[day] = data.get(day, "fill")
            out[name] = row
    return out

def _default_shift_times():
    return {
        "SS_weekday_close":  "13:30–22:00",
        "SS_weekday_open":   "08:00–16:30",
        "PS_weekday_close":  "11:30–20:00",
        "PS_weekday_open":   "08:00–16:30",
        "SS_saturday_open":  "09:00–17:30",
        "SS_saturday_close": "10:30–19:00",
        "SS_sunday_close":   "09:30–18:00",
    }

def _on_pto(name, d, pto_dict):
    return d in pto_dict.get(name, set())

def _stat_dates(norm_holidays):
    return {h["date"]: h["name"] for h in norm_holidays}


# =========================================================
# SS WEEKDAY CLOSING ENGINE
# =========================================================

def _gen_ss_weekday(mondays, rotation, pto, skip_start):
    """
    Each week one TL 'skips' (sits out). The remaining n-1 TLs cover Mon–Fri.
    Shavindri never closes on Monday.
    """
    n = len(rotation)
    if n == 0:
        return []
    skip_idx = skip_start % n
    results  = []

    for week_start in mondays:
        skip_tl    = rotation[skip_idx % n]
        candidates = [rotation[(skip_idx + 1 + i) % n] for i in range(n - 1)]

        # Shavindri can't close Monday — swap her to Tuesday position if she'd be first
        if len(candidates) >= 2 and candidates[0] == "Shavindri":
            candidates[0], candidates[1] = candidates[1], candidates[0]

        used  = set()
        fills = {}
        row   = {"week_start": week_start, "skip_tl": skip_tl, "fills": fills}

        for di, day in enumerate(["Mon", "Tue", "Wed", "Thu", "Fri"]):
            day_date = week_start + datetime.timedelta(days=di)
            assigned = None

            for tl in candidates:
                if tl in used:
                    continue
                if _on_pto(tl, day_date, pto):
                    continue
                if di == 0 and tl == "Shavindri":
                    continue
                assigned = tl
                break

            # Natural order substitution note
            if assigned:
                pos     = len(used)
                natural = candidates[pos] if pos < len(candidates) else None
                if natural and natural != assigned and _on_pto(natural, day_date, pto):
                    fills[day] = f"{natural} PTO → {assigned} fills"

            # Fallback 1: skip TL covers
            if assigned is None:
                if (skip_tl not in used
                        and not _on_pto(skip_tl, day_date, pto)
                        and not (di == 0 and skip_tl == "Shavindri")):
                    assigned          = skip_tl
                    fills[day]        = f"{skip_tl} fills (skip TL covers)"
                else:
                    # Fallback 2: anyone in the full rotation
                    for tl in rotation:
                        if (tl not in used
                                and not _on_pto(tl, day_date, pto)
                                and not (di == 0 and tl == "Shavindri")):
                            assigned   = tl
                            fills[day] = f"{tl} fills (emergency)"
                            break
                    if not assigned:
                        assigned   = "UNASSIGNED"
                        fills[day] = "No TL available"

            used.add(assigned)
            row[day] = assigned

        results.append(row)
        skip_idx = (skip_idx + 1) % n

    return results


# =========================================================
# SS WEEKEND ENGINES
# =========================================================

def _gen_ss_sunday(mondays, rotation, pto, sun_start):
    n = len(rotation)
    if n == 0:
        return {}
    idx    = sun_start % n
    result = {}
    for week_start in mondays:
        sun_date = week_start + datetime.timedelta(days=6)
        tl, tries = rotation[idx % n], 0
        while _on_pto(tl, sun_date, pto) and tries < n:
            idx   = (idx + 1) % n
            tl    = rotation[idx % n]
            tries += 1
        result[week_start] = tl
        idx = (idx + 1) % n
    return result

def _gen_ss_saturday(mondays, primary_tl, fill_rotation, pto):
    eligible = [t for t in fill_rotation if t != primary_tl]
    fill_idx = 0
    n        = len(eligible)
    result   = {}
    for week_start in mondays:
        sat_date = week_start + datetime.timedelta(days=5)
        if primary_tl and not _on_pto(primary_tl, sat_date, pto):
            result[week_start] = (primary_tl, False)
        else:
            found = False
            for _ in range(max(n, 1)):
                if n == 0:
                    break
                tl       = eligible[fill_idx % n]
                fill_idx = (fill_idx + 1) % n
                if not _on_pto(tl, sat_date, pto):
                    result[week_start] = (f"{tl} (fills)", True)
                    found              = True
                    break
            if not found:
                result[week_start] = ("UNASSIGNED", True)
    return result


# =========================================================
# PS ENGINE
# =========================================================

def _gen_ps(mondays, fri_rotation, weekend_tl, monday_thu_tl, pto, fri_start):
    n       = len(fri_rotation)
    fri_idx = fri_start % max(n, 1)
    results = []

    for week_start in mondays:
        fri_date  = week_start + datetime.timedelta(days=4)
        fri_tl    = fri_rotation[fri_idx % n] if n else "UNASSIGNED"
        fill_note = ""

        if n and _on_pto(fri_tl, fri_date, pto):
            original = fri_tl
            for offset in range(1, n + 1):
                alt = fri_rotation[(fri_idx + offset) % n]
                if not _on_pto(alt, fri_date, pto):
                    fri_tl    = alt
                    fill_note = f"{original} PTO → {alt} covers Friday"
                    break
            else:
                fri_tl    = "UNASSIGNED"
                fill_note = f"{original} PTO — no Friday cover"

        results.append({
            "week_start": week_start,
            "Mon":        monday_thu_tl,
            "Tue":        monday_thu_tl,
            "Wed":        monday_thu_tl,
            "Thu":        monday_thu_tl,
            "Fri":        fri_tl,
            "Sat":        f"{weekend_tl} (primary)" if weekend_tl else "",
            "Sun":        weekend_tl,
            "fill_note":  fill_note,
        })

        if n:
            fri_idx = (fri_idx + 1) % n

    return results


# =========================================================
# STAT HOLIDAY ASSIGNMENT
# =========================================================

def _assign_stats(norm_holidays, ss_sched, ps_sched, ss_sats, ss_suns, shift_times):
    ss_map = {r["week_start"]: r for r in ss_sched}
    ps_map = {r["week_start"]: r for r in ps_sched}
    recs   = []

    for h in norm_holidays:
        d    = h["date"]
        name = h["name"]
        dn   = _day_name(d)
        ws   = d - datetime.timedelta(days=d.weekday())
        ss   = ss_map.get(ws, {})
        ps   = ps_map.get(ws, {})

        if dn in ("Mon", "Tue", "Wed", "Thu", "Fri"):
            ss_closer  = ss.get(dn, "TBD")
            ps_opener  = ps.get("Mon", "TBD") if dn in ("Mon","Tue","Wed","Thu") else ps.get("Fri", "TBD")
            ss_shift   = shift_times.get("SS_weekday_close", "")
        elif dn == "Sat":
            sat_tl, _  = ss_sats.get(ws, ("TBD", False))
            ss_closer  = sat_tl
            ps_opener  = ps.get("Sat", "TBD")
            ss_shift   = shift_times.get("SS_saturday_close", "")
        else:
            ss_closer  = ss_suns.get(ws, "TBD")
            ps_opener  = ps.get("Sun", "TBD")
            ss_shift   = shift_times.get("SS_sunday_close", "")

        recs.append({
            "Date":      str(d),
            "Day":       dn,
            "Holiday":   name,
            "SS Closer": ss_closer,
            "SS Shift":  ss_shift,
            "PS Opener": ps_opener,
            "PS Shift":  shift_times.get("PS_weekday_open", ""),
        })

    return recs


# =========================================================
# SHIFT COUNTS (for summary tab)
# =========================================================

def _build_shift_counts(ss_sched, ps_sched, ss_sats, ss_suns, all_tls):
    counts = {n: defaultdict(int) for n in all_tls}

    for row in ss_sched:
        for day in ("Mon", "Tue", "Wed", "Thu", "Fri"):
            tl = row.get(day, "")
            if tl and tl in counts:
                counts[tl][day]   += 1
                counts[tl]["Total"] += 1

    for week_start, (sat_tl, _) in ss_sats.items():
        tl = sat_tl.replace(" (fills)", "").strip()
        if tl and tl in counts:
            counts[tl]["Sat"]   += 1
            counts[tl]["Total"] += 1

    for week_start, sun_tl in ss_suns.items():
        if sun_tl and sun_tl in counts:
            counts[sun_tl]["Sun"]   += 1
            counts[sun_tl]["Total"] += 1

    for row in ps_sched:
        for day in ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"):
            tl = row.get(day, "").replace(" (primary)", "").strip()
            if tl and tl in counts:
                counts[tl][day]   += 1
                counts[tl]["Total"] += 1

    return counts


# =========================================================
# TEXT REPORT
# =========================================================

def _build_text_report(ss_sched, ps_sched, ss_sats, ss_suns, stat_recs, start_dt, end_dt):
    lines = []
    lines.append("=" * 65)
    lines.append("  WFM TEAM LEAD ROTATION SCHEDULE")
    lines.append(f"  Generated : {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}")
    lines.append(f"  Period    : {start_dt}  →  {end_dt}")
    lines.append("=" * 65)

    lines.append("\n── SS CLOSING ROTATION ──")
    header = f"{'Date':<10} {'Mon':<12} {'Tue':<12} {'Wed':<12} {'Thu':<12} {'Fri':<12} {'Sat':<14} {'Sun/LOD':<12}"
    lines.append(header)
    for row in ss_sched:
        ws  = row["week_start"]
        sat = ss_sats.get(ws, ("", False))[0]
        sun = ss_suns.get(ws, "")
        fills = row.get("fills", {})
        line = (f"{ws.strftime('%d-%b'):<10} "
                f"{row.get('Mon',''):<12} {row.get('Tue',''):<12} "
                f"{row.get('Wed',''):<12} {row.get('Thu',''):<12} "
                f"{row.get('Fri',''):<12} {sat:<14} {sun:<12}")
        lines.append(line)
        for day, note in fills.items():
            lines.append(f"           ⚠ {day}: {note}")

    lines.append("\n── PS SCHEDULE ──")
    for row in ps_sched:
        ws   = row["week_start"]
        note = f"  [{row['fill_note']}]" if row.get("fill_note") else ""
        lines.append(
            f"{ws.strftime('%d-%b'):<10} Mon–Thu: {row.get('Mon','')}  "
            f"Fri: {row.get('Fri','')}  Sat/Sun: {row.get('Sun','')}{note}"
        )

    if stat_recs:
        lines.append("\n── STAT HOLIDAYS ──")
        lines.append(f"{'Date':<12} {'Day':<5} {'Holiday':<22} {'SS Closer':<14} {'SS Shift':<14} {'PS Opener':<12} {'PS Shift'}")
        for r in stat_recs:
            lines.append(
                f"{r['Date']:<12} {r['Day']:<5} {r['Holiday']:<22} "
                f"{r['SS Closer']:<14} {r['SS Shift']:<14} {r['PS Opener']:<12} {r['PS Shift']}"
            )

    lines.append("\n" + "=" * 65)
    return "\n".join(lines)


# =========================================================
# HTML REPORT
# =========================================================

def _build_html_report(ss_sched, ps_sched, ss_sats, ss_suns, stat_recs, start_dt, end_dt, shift_times):
    rows_ss = []
    for row in ss_sched:
        ws    = row["week_start"]
        sat   = ss_sats.get(ws, ("", False))[0]
        sun   = ss_suns.get(ws, "")
        fills = row.get("fills", {})
        notes = " | ".join(f"{d}: {m}" for d, m in fills.items())
        rows_ss.append(
            f"<tr>"
            f"<td>{ws.strftime('%d-%b')}</td>"
            f"<td>{_cell(row.get('Mon',''), fills.get('Mon'))}</td>"
            f"<td>{_cell(row.get('Tue',''), fills.get('Tue'))}</td>"
            f"<td>{_cell(row.get('Wed',''), fills.get('Wed'))}</td>"
            f"<td>{_cell(row.get('Thu',''), fills.get('Thu'))}</td>"
            f"<td>{_cell(row.get('Fri',''), fills.get('Fri'))}</td>"
            f"<td>{_cell(sat, None, 'fills' in sat)}</td>"
            f"<td>{sun}</td>"
            f"<td style='font-size:.8rem;color:#7B4F00'>{notes}</td>"
            f"</tr>"
        )

    rows_ps = []
    for row in ps_sched:
        ws   = row["week_start"]
        note = row.get("fill_note", "")
        rows_ps.append(
            f"<tr>"
            f"<td>{ws.strftime('%d-%b')}</td>"
            f"<td>{row.get('Mon','')}</td>"
            f"<td>{row.get('Tue','')}</td>"
            f"<td>{row.get('Wed','')}</td>"
            f"<td>{row.get('Thu','')}</td>"
            f"<td>{_cell(row.get('Fri',''), note if note else None)}</td>"
            f"<td>{row.get('Sat','')}</td>"
            f"<td>{row.get('Sun','')}</td>"
            f"<td style='font-size:.8rem;color:#7B4F00'>{note}</td>"
            f"</tr>"
        )

    rows_stat = "".join(
        f"<tr><td>{r['Date']}</td><td>{r['Day']}</td><td>{r['Holiday']}</td>"
        f"<td><strong>{r['SS Closer']}</strong></td><td>{r['SS Shift']}</td>"
        f"<td><strong>{r['PS Opener']}</strong></td><td>{r['PS Shift']}</td></tr>"
        for r in stat_recs
    )

    thead = "<tr><th>Date</th><th>Mon</th><th>Tue</th><th>Wed</th><th>Thu</th><th>Fri</th><th>Sat</th><th>Sun/LOD</th><th>Notes</th></tr>"

    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8"/>
  <title>Schedule Report</title>
  <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.2/dist/css/bootstrap.min.css" rel="stylesheet">
</head>
<body class="p-4" style="background:#f8f9fa;">
<div class="container-fluid">
  <div class="mb-3">
    <h2 class="mb-0">📅 TL Rotation Schedule</h2>
    <div class="text-muted small">
      Generated {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')} &nbsp;|&nbsp;
      Period: {start_dt} → {end_dt}
    </div>
  </div>

  <h5 class="mt-4">📋 SS Closing Rotation</h5>
  <div class="table-responsive">
    <table class="table table-sm table-bordered table-hover" style="font-size:.85rem">
      <thead class="table-dark">{thead}</thead>
      <tbody>{"".join(rows_ss)}</tbody>
    </table>
  </div>

  <h5 class="mt-4">📋 PS Schedule</h5>
  <div class="table-responsive">
    <table class="table table-sm table-bordered table-hover" style="font-size:.85rem">
      <thead class="table-success">{thead}</thead>
      <tbody>{"".join(rows_ps)}</tbody>
    </table>
  </div>

  {"" if not stat_recs else f'''
  <h5 class="mt-4">⭐ Stat Holidays</h5>
  <div class="table-responsive">
    <table class="table table-sm table-bordered" style="font-size:.85rem">
      <thead class="table-danger">
        <tr><th>Date</th><th>Day</th><th>Holiday</th>
            <th>SS Closer</th><th>SS Shift</th>
            <th>PS Opener</th><th>PS Shift</th></tr>
      </thead>
      <tbody>{rows_stat}</tbody>
    </table>
  </div>'''}

</div>
</body>
</html>"""

def _cell(value, note=None, is_fill=False):
    if is_fill or (note and "fills" in str(note)):
        return f'<span style="background:#FFD966;padding:1px 4px;border-radius:3px">{value}</span>'
    return value or ""


# =========================================================
# EXCEL EXPORT
# =========================================================

def _build_excel(ss_sched, ps_sched, ss_sats, ss_suns, stat_recs,
                 norm_pto, norm_accom, norm_holidays, all_tls, shift_times):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        log.error("openpyxl not installed — Excel export unavailable")
        return None

    wb  = Workbook()
    thn = Side(style="thin", color="B4B4B4")
    bdr = Border(left=thn, right=thn, top=thn, bottom=thn)

    def _fill(h):
        return PatternFill("solid", fgColor=h.lstrip("#"))

    def _write_sheet(ws, headers, data_rows, hdr_color, alt_color, kw=None):
        """Write a header row + data rows with alternating fill."""
        for ci, col in enumerate(headers, 1):
            c = ws.cell(1, ci, col)
            c.fill      = _fill(hdr_color)
            c.font      = Font(bold=True, color="FFFFFF", size=11, name="Arial")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border    = bdr
        ws.row_dimensions[1].height = 25

        for ri, row in enumerate(data_rows, 2):
            bg = alt_color if ri % 2 == 0 else "FFFFFF"
            for ci, val in enumerate(row, 1):
                cell_val = "" if val is None else str(val)
                c = ws.cell(ri, ci, cell_val)
                c.border    = bdr
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
                c.font      = Font(name="Arial", size=10)
                cell_bg = bg
                if kw and kw in cell_val:
                    cell_bg = "FFD966"
                    c.font  = Font(bold=True, color="9C5700", name="Arial", size=10)
                elif ci == len(headers) and cell_val.strip():  # Notes col
                    cell_bg = "FFF8E1"
                    c.font  = Font(color="7B4F00", italic=True, name="Arial", size=9)
                    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                c.fill = _fill(cell_bg)
            ws.row_dimensions[ri].height = 20

    def _col_widths(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ── SS Rotation sheet ──
    ws_ss = wb.active
    ws_ss.title = "SS Rotation"
    ss_headers = ["Date", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday / LOD", "Notes"]
    ss_rows = []
    stat_dates = _stat_dates(norm_holidays)
    for row in ss_sched:
        ws_d  = row["week_start"]
        sat   = ss_sats.get(ws_d, ("", False))[0]
        sun   = ss_suns.get(ws_d, "")
        fills = row.get("fills", {})
        notes = [f"{d}: {m}" for d, m in fills.items()]
        for d, hname in stat_dates.items():
            if ws_d <= d <= ws_d + datetime.timedelta(days=6):
                dn = _day_name(d)
                if dn in ("Mon","Tue","Wed","Thu","Fri"):
                    notes.append(f"⭐ {hname} ({dn}) = {row.get(dn,'?')} closes")
                elif dn == "Sat":
                    notes.append(f"⭐ {hname} (Sat) = {sat} closes")
        ss_rows.append([
            ws_d.strftime("%d-%b"),
            row.get("Mon",""), row.get("Tue",""), row.get("Wed",""),
            row.get("Thu",""), row.get("Fri",""), sat, sun,
            " | ".join(notes),
        ])
    _write_sheet(ws_ss, ss_headers, ss_rows, "1F4973", "EBF1FC", kw="(fills)")
    _col_widths(ws_ss, [10,13,13,13,13,13,16,16,60])
    ws_ss.freeze_panes = "A2"

    # ── PS Rotation sheet ──
    ws_ps = wb.create_sheet("PS Rotation")
    ps_headers = ["Date", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday / LOD", "Notes"]
    ps_rows_out = []
    for row in ps_sched:
        ws_d  = row["week_start"]
        notes = [row["fill_note"]] if row.get("fill_note") else []
        for d, hname in stat_dates.items():
            if ws_d <= d <= ws_d + datetime.timedelta(days=6):
                dn     = _day_name(d)
                opener = row.get("Mon") if dn in ("Mon","Tue","Wed","Thu") else (row.get("Fri") if dn == "Fri" else row.get("Sun","Aman"))
                notes.append(f"⭐ {hname} ({dn}) = {opener} opens")
        ps_rows_out.append([
            ws_d.strftime("%d-%b"),
            row.get("Mon",""), row.get("Tue",""), row.get("Wed",""),
            row.get("Thu",""), row.get("Fri",""),
            row.get("Sat",""), row.get("Sun",""),
            " | ".join(notes),
        ])
    _write_sheet(ws_ps, ps_headers, ps_rows_out, "548235", "EBF4E6")
    _col_widths(ws_ps, [10,13,13,13,13,13,16,16,60])
    ws_ps.freeze_panes = "A2"

    # ── Stat Holidays sheet ──
    ws_st = wb.create_sheet("Stat Holidays")
    _write_sheet(ws_st,
        ["Date","Day","Holiday","SS Closer","SS Shift","PS Opener","PS Shift"],
        [[r["Date"],r["Day"],r["Holiday"],r["SS Closer"],r["SS Shift"],r["PS Opener"],r["PS Shift"]]
         for r in stat_recs],
        "7B2C2C", "FFEEEE")
    _col_widths(ws_st, [14,6,22,16,14,16,14])

    # ── PTO sheet ──
    ws_pt = wb.create_sheet("PTO")
    pto_rows = []
    for name, dates in norm_pto.items():
        if not dates:
            continue
        sorted_dates = sorted(dates)
        # Collapse contiguous ranges
        ranges, start_r = [], sorted_dates[0]
        prev = sorted_dates[0]
        for d in sorted_dates[1:]:
            if (d - prev).days > 1:
                ranges.append((start_r, prev))
                start_r = d
            prev = d
        ranges.append((start_r, prev))
        for s, e in ranges:
            pto_rows.append([name, all_tls.get(name,""), str(s), str(e)])
    _write_sheet(ws_pt, ["Name","Group","Start","End"], pto_rows, "6B4C9A", "F5F0FF")
    _col_widths(ws_pt, [16,8,14,14])

    # ── Accommodations sheet ──
    ws_ac = wb.create_sheet("Accommodations")
    accom_rows = [
        [name, data.get("group",""),
         data.get("Mon",""), data.get("Tue",""), data.get("Wed",""),
         data.get("Thu",""), data.get("Fri",""), data.get("Sat",""), data.get("Sun",""),
         data.get("notes","")]
        for name, data in norm_accom.items()
    ]
    _write_sheet(ws_ac,
        ["Name","Group","Mon","Tue","Wed","Thu","Fri","Sat","Sun","Notes"],
        accom_rows, "2E4057", "F0F4F8")
    _col_widths(ws_ac, [14,6,8,8,8,8,8,8,8,50])

    # ── Summary sheet ──
    ws_sm = wb.create_sheet("Summary")
    counts = _build_shift_counts(ss_sched, ps_sched, ss_sats, ss_suns, all_tls)
    summary_rows = [
        [name, all_tls.get(name,""),
         counts[name]["Mon"], counts[name]["Tue"], counts[name]["Wed"],
         counts[name]["Thu"], counts[name]["Fri"], counts[name]["Sat"],
         counts[name]["Sun"], counts[name]["Total"]]
        for name in all_tls
    ]
    _write_sheet(ws_sm,
        ["Name","Group","Mon","Tue","Wed","Thu","Fri","Sat","Sun","Total"],
        summary_rows, "4A4A4A", "F5F5F5")
    _col_widths(ws_sm, [14,6,9,9,9,9,9,9,9,9])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# =========================================================
# STANDALONE TEST
# =========================================================

if __name__ == "__main__":
    result = build_schedule(
        ss_list=["Alana","Kristin","Shavindri","Odette","Salomon","Cushana","Lucas","Adeyinka"],
        ps_list=["Jodi","David","Edena","Bryan","Julianne","Aman"],
        ss_closing_rotation=["Odette","Salomon","Shavindri","Cushana","Lucas","Adeyinka"],
        ss_sunday_rotation=["Cushana","Lucas","Adeyinka","Odette","Salomon"],
        ps_friday_rotation=["Bryan","Julianne","Edena"],
        ps_weekend_tl="Aman",
        ps_monday_thursday_tl="David",
        ss_saturday_primary_tl="Shavindri",
        ss_skip_start=2,
        ps_fri_start=0,
        ss_sun_start=0,
        schedule_start="2026-06-29",
        schedule_end="2026-12-28",
        pto=[
            {"name":"Shavindri","start":"2026-07-15","end":"2026-07-23"},
            {"name":"Lucas",    "start":"2026-08-04","end":"2026-08-06"},
        ],
        holidays=[
            {"date":"2026-07-01","name":"Canada Day"},
            {"date":"2026-08-03","name":"Civic Holiday"},
            {"date":"2026-09-07","name":"Labour Day"},
            {"date":"2026-10-12","name":"Thanksgiving"},
            {"date":"2026-12-25","name":"Christmas Day"},
            {"date":"2026-12-26","name":"Boxing Day"},
        ],
    )
    print(result["text_report"])
    if result["excel_bytes"]:
        with open("TL_Rotation_Schedule_test.xlsx", "wb") as f:
            f.write(result["excel_bytes"].read())
        print("Excel saved: TL_Rotation_Schedule_test.xlsx")
