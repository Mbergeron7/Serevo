"""
routes/rosters.py — Roster Rotation blueprint
===============================================
CRUD for roster groups, members, slots, and roster generation/publish.
"""

import csv
import io
import json
import logging
import datetime

from flask import Blueprint, render_template, request, jsonify, Response
from app.auth import login_required, admin_required, get_current_user
from app.models import (db, RosterGroup, RosterMember, RosterSlot,
                         RosterRotationState, Employee, PlanningUnit, Schedule,
                         SpecialDay)
from app.scheduling.roster_engine import (
    generate_roster, publish_roster, clear_roster_drafts, get_roster_calendar,
)

log = logging.getLogger("serevo.rosters")

rosters_bp = Blueprint("rosters", __name__, url_prefix="/scheduling/rosters")


# ── Page ─────────────────────────────────────────────────────

@rosters_bp.route("/")
@rosters_bp.route("")
@login_required
def index():
    user = get_current_user()
    return render_template("rosters.html", user=user)


# ── Planning Units (for dropdowns) ───────────────────────────

@rosters_bp.route("/api/planning-units")
@login_required
def api_planning_units():
    pus = PlanningUnit.get_all_active()
    return jsonify([{"id": p.id, "name": p.name} for p in pus])


# ── Employees (for member picker) ────────────────────────────

@rosters_bp.route("/api/employees")
@login_required
def api_employees():
    pu_id = request.args.get("planning_unit_id", type=int)
    q = Employee.query.filter_by(status="Active")
    if pu_id:
        q = q.filter_by(planning_unit_id=pu_id)
    emps = q.order_by(Employee.last_name, Employee.first_name).all()
    return jsonify([{"id": e.id, "employee_id": e.employee_id,
                     "name": e.full_name} for e in emps])


# ── Groups CRUD ──────────────────────────────────────────────

@rosters_bp.route("/api/groups")
@login_required
def api_groups():
    groups = RosterGroup.query.filter_by(active=True).order_by(RosterGroup.name).all()
    return jsonify([g.to_dict() for g in groups])


@rosters_bp.route("/api/groups", methods=["POST"])
@login_required
def api_create_group():
    d = request.get_json(force=True)
    g = RosterGroup(
        planning_unit_id=d["planning_unit_id"],
        name=d["name"],
        closers_required_weekday=d.get("closers_required_weekday", 1),
        closers_required_weekend=d.get("closers_required_weekend", 1),
        weekend_mode=d.get("weekend_mode", "rotate"),
        stat_role=d.get("stat_role"),
        sunday_non_closer=d.get("sunday_non_closer", False),
        weekend_close_week_parity=d.get("weekend_close_week_parity") or None,
    )
    db.session.add(g)
    db.session.commit()
    return jsonify(g.to_dict()), 201


@rosters_bp.route("/api/groups/<int:gid>", methods=["PUT"])
@login_required
def api_update_group(gid):
    g = RosterGroup.query.get_or_404(gid)
    d = request.get_json(force=True)
    for k in ("name", "planning_unit_id", "closers_required_weekday",
              "closers_required_weekend", "weekend_mode", "stat_role",
              "sunday_non_closer", "weekend_close_week_parity"):
        if k in d:
            val = d[k]
            if k == "weekend_close_week_parity":
                val = val or None  # empty string → null
            setattr(g, k, val)
    db.session.commit()
    return jsonify(g.to_dict())


@rosters_bp.route("/api/groups/<int:gid>", methods=["DELETE"])
@login_required
def api_delete_group(gid):
    g = RosterGroup.query.get_or_404(gid)
    g.active = False
    db.session.commit()
    return jsonify({"ok": True})


# ── Members CRUD ─────────────────────────────────────────────

@rosters_bp.route("/api/groups/<int:gid>/members", methods=["POST"])
@login_required
def api_add_member(gid):
    g = RosterGroup.query.get_or_404(gid)
    d = request.get_json(force=True)
    m = RosterMember(
        roster_group_id=g.id,
        employee_id=d["employee_id"],
        is_display_only=d.get("is_display_only", False),
        is_weekend_primary=d.get("is_weekend_primary", False),
        is_fill_in=d.get("is_fill_in", False),
        weekend_primary_days=json.dumps(d["weekend_primary_days"]) if d.get("weekend_primary_days") else None,
        max_closings_per_week=d.get("max_closings_per_week", 1),
        max_openings_per_week=d.get("max_openings_per_week"),
        max_mids_per_week=d.get("max_mids_per_week"),
        excluded_closing_days=json.dumps(d["excluded_closing_days"]) if d.get("excluded_closing_days") else None,
        static_days_json=json.dumps(d["static_days_json"]) if d.get("static_days_json") else None,
        accommodation_start=_parse_time(d.get("accommodation_start")),
        accommodation_end=_parse_time(d.get("accommodation_end")),
        accommodation_no_close=d.get("accommodation_no_close", False),
        dedicated_closer=d.get("dedicated_closer", False),
        dedicated_opener=d.get("dedicated_opener", False),
        exclude_from_saturday=d.get("exclude_from_saturday", d.get("exclude_from_weekends", False)),
        exclude_from_sunday=d.get("exclude_from_sunday", d.get("exclude_from_weekends", False)),
        exclude_from_opening=d.get("exclude_from_opening", False),
        exclude_from_closing=d.get("exclude_from_closing", False),
        excluded_opening_days=json.dumps(d["excluded_opening_days"]) if d.get("excluded_opening_days") else None,
        allow_split_shift=d.get("allow_split_shift", False),
        split_shift_min_gap=d.get("split_shift_min_gap"),
        consistent_start=d.get("consistent_start", False),
        max_consecutive_days=d.get("max_consecutive_days"),
        preferred_days_off=json.dumps(d["preferred_days_off"]) if d.get("preferred_days_off") else None,
        min_hours_per_week=d.get("min_hours_per_week"),
        max_hours_per_week=d.get("max_hours_per_week"),
        pair_with_member_id=d.get("pair_with_member_id"),
        avoid_member_id=d.get("avoid_member_id"),
        work_days=json.dumps(d["work_days"]) if d.get("work_days") else None,
        contracted_days=d.get("contracted_days", 5),
        contracted_hours=d.get("contracted_hours"),
        rotation_order=d.get("rotation_order", 0),
    )
    db.session.add(m)
    db.session.commit()
    return jsonify(m.to_dict()), 201


@rosters_bp.route("/api/members/<int:mid>", methods=["PUT"])
@login_required
def api_update_member(mid):
    m = RosterMember.query.get_or_404(mid)
    d = request.get_json(force=True)
    for k in ("is_display_only", "is_weekend_primary", "is_fill_in",
              "max_closings_per_week", "max_openings_per_week", "max_mids_per_week",
              "accommodation_no_close", "rotation_order",
              "contracted_days", "contracted_hours",
              "dedicated_closer", "dedicated_opener",
              "exclude_from_saturday", "exclude_from_sunday",
              "exclude_from_opening", "exclude_from_closing",
              "allow_split_shift", "split_shift_min_gap",
              "consistent_start",
              "max_consecutive_days", "min_hours_per_week", "max_hours_per_week",
              "pair_with_member_id", "avoid_member_id"):
        if k in d:
            setattr(m, k, d[k])
    for k in ("weekend_primary_days", "excluded_closing_days", "excluded_opening_days",
              "static_days_json", "work_days", "preferred_days_off"):
        if k in d:
            setattr(m, k, json.dumps(d[k]) if d[k] else None)
    if "accommodation_start" in d:
        m.accommodation_start = _parse_time(d["accommodation_start"])
    if "accommodation_end" in d:
        m.accommodation_end = _parse_time(d["accommodation_end"])
    db.session.commit()
    return jsonify(m.to_dict())


@rosters_bp.route("/api/members/<int:mid>", methods=["DELETE"])
@login_required
def api_delete_member(mid):
    m = RosterMember.query.get_or_404(mid)
    m.active = False
    db.session.commit()
    return jsonify({"ok": True})


# ── Slots CRUD ───────────────────────────────────────────────

@rosters_bp.route("/api/groups/<int:gid>/slots", methods=["POST"])
@login_required
def api_add_slot(gid):
    g = RosterGroup.query.get_or_404(gid)
    d = request.get_json(force=True)
    s = RosterSlot(
        roster_group_id=g.id,
        day_type=d["day_type"],
        role=d["role"],
        label=d.get("label"),
        shift_start=_parse_time(d["shift_start"]),
        shift_end=_parse_time(d["shift_end"]),
        sort_order=d.get("sort_order", 0),
    )
    db.session.add(s)
    db.session.commit()
    return jsonify(s.to_dict()), 201


@rosters_bp.route("/api/slots/<int:sid>", methods=["PUT"])
@login_required
def api_update_slot(sid):
    s = RosterSlot.query.get_or_404(sid)
    d = request.get_json(force=True)
    if "day_type" in d:
        s.day_type = d["day_type"]
    if "role" in d:
        s.role = d["role"]
    if "label" in d:
        s.label = d["label"]
    if "shift_start" in d:
        s.shift_start = _parse_time(d["shift_start"])
    if "shift_end" in d:
        s.shift_end = _parse_time(d["shift_end"])
    if "sort_order" in d:
        s.sort_order = d.get("sort_order", 0)
    db.session.commit()
    return jsonify(s.to_dict())


@rosters_bp.route("/api/slots/<int:sid>", methods=["DELETE"])
@login_required
def api_delete_slot(sid):
    s = RosterSlot.query.get_or_404(sid)
    db.session.delete(s)
    db.session.commit()
    return jsonify({"ok": True})


# ── Bulk slot save (save all slots for a group at once) ──────

@rosters_bp.route("/api/groups/<int:gid>/slots/bulk", methods=["POST"])
@login_required
def api_bulk_save_slots(gid):
    g = RosterGroup.query.get_or_404(gid)
    data = request.get_json(force=True)
    slots = data.get("slots", [])
    # Remove existing slots for this group
    RosterSlot.query.filter_by(roster_group_id=g.id).delete()
    for s in slots:
        db.session.add(RosterSlot(
            roster_group_id=g.id,
            day_type=s["day_type"],
            role=s["role"],
            label=s.get("label"),
            shift_start=_parse_time(s["shift_start"]),
            shift_end=_parse_time(s["shift_end"]),
            sort_order=s.get("sort_order", 0),
        ))
    db.session.commit()
    return jsonify({"ok": True, "count": len(slots)})


# ── Generation / Publish / Clear ─────────────────────────────

@rosters_bp.route("/api/generate", methods=["POST"])
@login_required
def api_generate():
    d = request.get_json(force=True)
    group_id = d.get("group_id")
    start = _parse_date(d.get("start_date"))
    end = _parse_date(d.get("end_date"))
    if not group_id or not start or not end:
        return jsonify({"error": "group_id, start_date, and end_date required"}), 400
    if end < start:
        return jsonify({"error": "end_date must be >= start_date"}), 400
    try:
        # Ensure a clean session (previous timed-out request may have
        # left the connection in a bad state on this worker).
        db.session.rollback()
        result = generate_roster(group_id, start, end)
    except Exception as e:
        db.session.rollback()
        log.exception("Roster generation failed: %s", e)
        return jsonify({"error": f"Generation failed: {e}"}), 500
    if "error" in result:
        return jsonify(result), 400
    return jsonify(result)


@rosters_bp.route("/api/publish", methods=["POST"])
@login_required
def api_publish():
    d = request.get_json(force=True)
    group_id = d.get("group_id")
    start = _parse_date(d.get("start_date"))
    end = _parse_date(d.get("end_date"))
    if not group_id or not start or not end:
        return jsonify({"error": "group_id, start_date, and end_date required"}), 400
    result = publish_roster(group_id, start, end)
    if "error" in result:
        return jsonify(result), 400
    return jsonify(result)


@rosters_bp.route("/api/debug-groups")
@login_required
def api_debug_groups():
    """List all roster groups with their IDs."""
    groups = RosterGroup.query.all()
    return jsonify([{"id": g.id, "name": g.name} for g in groups])


@rosters_bp.route("/api/debug-slots/<int:gid>")
@login_required
def api_debug_slots(gid):
    """Debug endpoint: dump what the engine sees for a group."""
    g = RosterGroup.query.get_or_404(gid)
    members = [m for m in g.members if m.active and not m.is_display_only]
    slots = []
    for s in g.slots:
        slots.append({
            "id": s.id, "day_type": s.day_type, "role": s.role,
            "label": s.label, "shift_start": str(s.shift_start),
            "shift_end": str(s.shift_end), "sort_order": s.sort_order,
        })
    mbrs = []
    for m in members:
        emp = m.employee
        name = f"{emp.first_name} {emp.last_name}" if emp else f"emp#{m.employee_id}"
        mbrs.append({
            "id": m.id, "name": name,
            "dedicated_closer": m.dedicated_closer,
            "dedicated_opener": m.dedicated_opener,
            "exclude_from_closing": m.exclude_from_closing,
            "exclude_from_opening": m.exclude_from_opening,
            "is_weekend_primary": m.is_weekend_primary,
            "weekend_primary_days": m.weekend_primary_days,
            "work_days": m.work_days,
            "excluded_closing_days": m.excluded_closing_days,
            "accommodation_no_close": m.accommodation_no_close,
            "max_closings_per_week": m.max_closings_per_week,
            "rotation_order": m.rotation_order,
            "is_fill_in": m.is_fill_in,
            "is_display_only": m.is_display_only,
            "static_days_json": m.static_days_json,
        })
    return jsonify({
        "group": {"id": g.id, "name": g.name,
                  "closers_required_weekday": g.closers_required_weekday,
                  "closers_required_weekend": g.closers_required_weekend,
                  "weekend_mode": g.weekend_mode,
                  "sunday_non_closer": g.sunday_non_closer,
                  "weekend_close_week_parity": g.weekend_close_week_parity},
        "slots": slots,
        "members": mbrs,
    })


@rosters_bp.route("/api/clear", methods=["POST"])
@login_required
def api_clear():
    d = request.get_json(force=True)
    group_id = d.get("group_id")
    start = _parse_date(d.get("start_date"))
    end = _parse_date(d.get("end_date"))
    if not group_id or not start or not end:
        return jsonify({"error": "group_id, start_date, and end_date required"}), 400
    result = clear_roster_drafts(group_id, start, end)
    if "error" in result:
        return jsonify(result), 400
    return jsonify(result)


@rosters_bp.route("/api/calendar")
@login_required
def api_calendar():
    group_id = request.args.get("group_id", type=int)
    start = _parse_date(request.args.get("start_date"))
    end = _parse_date(request.args.get("end_date"))
    if not group_id or not start or not end:
        return jsonify({"error": "group_id, start_date, and end_date required"}), 400
    result = get_roster_calendar(group_id, start, end)
    if "error" in result:
        return jsonify(result), 400
    return jsonify(result)


# ── Export Groups & Members CSV ─────────────────────────────

@rosters_bp.route("/api/export-groups-csv")
@login_required
def api_export_groups_csv():
    """Export all roster groups and their members as CSV."""
    groups = RosterGroup.query.order_by(RosterGroup.name).all()
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([
        "Group", "Member", "Employee ID", "Role", "Rotation Order",
        "Work Days", "Max Closings/Week", "Weekend Primary", "Weekend Days",
        "Fill-In Eligible", "Display Only", "No Closing Shifts",
        "Dedicated Closer", "Dedicated Opener",
        "Exclude Opening", "Exclude Closing",
    ])
    for g in groups:
        members = RosterMember.query.filter_by(roster_group_id=g.id).order_by(
            RosterMember.rotation_order
        ).all()
        for m in members:
            emp = Employee.query.get(m.employee_id)
            name = emp.full_name if emp else f"ID {m.employee_id}"
            eid = emp.employee_id if emp else ""
            work_days = ""
            try:
                wd = json.loads(m.work_days) if m.work_days else None
                if wd:
                    work_days = ", ".join(wd)
            except Exception:
                pass
            wpd = ""
            try:
                wp = json.loads(m.weekend_primary_days) if m.weekend_primary_days else None
                if wp:
                    wpd = ", ".join(wp)
            except Exception:
                pass
            writer.writerow([
                g.name, name, eid, m.role or "",
                m.rotation_order or "", work_days,
                m.max_closings_per_week if m.max_closings_per_week is not None else "",
                "Yes" if m.is_weekend_primary else "",
                wpd,
                "Yes" if m.is_fill_in else "",
                "Yes" if m.is_display_only else "",
                "Yes" if m.accommodation_no_close else "",
                "Yes" if m.dedicated_closer else "",
                "Yes" if m.dedicated_opener else "",
                "Yes" if m.exclude_from_opening else "",
                "Yes" if m.exclude_from_closing else "",
            ])

    return Response(
        buf.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=roster_groups_members.csv"},
    )


# ── Export CSV ──────────────────────────────────────────────

@rosters_bp.route("/api/export-csv", methods=["POST"])
@login_required
def api_export_csv():
    d = request.get_json(force=True)
    group_id = d.get("group_id")
    start = _parse_date(d.get("start_date"))
    end = _parse_date(d.get("end_date"))
    if not group_id or not start or not end:
        return jsonify({"error": "group_id, start_date, and end_date required"}), 400

    result = get_roster_calendar(group_id, start, end)
    rows = result.get("rows", [])
    group = RosterGroup.query.get(group_id)
    group_name = group.name if group else "roster"

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Date", "Employee", "Role", "Shift Start", "Shift End", "Status"])
    for r in rows:
        writer.writerow([r["date"], r["employee"], r["type"], r["start"], r["end"], r["status"]])

    # ── Assignment Summary section ──
    writer.writerow([])
    writer.writerow(["=== Assignment Summary ==="])
    writer.writerow([])

    day_names = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    emp_sum = {}
    for r in rows:
        name = r.get("employee", "Unknown")
        if name not in emp_sum:
            emp_sum[name] = {"days": [0]*7, "sat": 0, "sun": 0,
                             "closings": 0, "openings": 0,
                             "close_days": [0]*7, "open_days": [0]*7}
        dow = None
        try:
            dt = datetime.date.fromisoformat(r["date"])
            dow = dt.weekday()  # 0=Mon
            emp_sum[name]["days"][dow] += 1
            if dow == 5:
                emp_sum[name]["sat"] += 1
            elif dow == 6:
                emp_sum[name]["sun"] += 1
        except (ValueError, AttributeError):
            pass
        role = (r.get("type") or "").lower()
        if role in ("closer", "close"):
            emp_sum[name]["closings"] += 1
            if dow is not None:
                emp_sum[name]["close_days"][dow] += 1
        if role in ("opener", "open"):
            emp_sum[name]["openings"] += 1
            if dow is not None:
                emp_sum[name]["open_days"][dow] += 1

    # Day-of-week counts
    writer.writerow(["Day-of-Week Counts"])
    writer.writerow(["Employee"] + day_names + ["Total"])
    for name in sorted(emp_sum):
        s = emp_sum[name]
        writer.writerow([name] + s["days"] + [sum(s["days"])])

    writer.writerow([])
    writer.writerow(["Weekend Counts"])
    writer.writerow(["Employee", "Saturday", "Sunday", "Total"])
    for name in sorted(emp_sum):
        s = emp_sum[name]
        writer.writerow([name, s["sat"], s["sun"], s["sat"] + s["sun"]])

    writer.writerow([])
    close_hdr = ["Employee", "Total"] + [f"Close {d}" for d in day_names]
    writer.writerow(["Closings Breakdown"])
    writer.writerow(close_hdr)
    for name in sorted(emp_sum):
        s = emp_sum[name]
        writer.writerow([name, s["closings"]] + s["close_days"])

    writer.writerow([])
    open_hdr = ["Employee", "Total"] + [f"Open {d}" for d in day_names]
    writer.writerow(["Openings Breakdown"])
    writer.writerow(open_hdr)
    for name in sorted(emp_sum):
        s = emp_sum[name]
        writer.writerow([name, s["openings"]] + s["open_days"])

    output = buf.getvalue()
    safe_name = group_name.replace(" ", "_")
    filename = f"roster_{safe_name}_{start}_{end}.csv"
    return Response(
        output,
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@rosters_bp.route("/api/export-xlsx", methods=["POST"])
@login_required
def api_export_xlsx():
    d = request.get_json(force=True)
    group_id = d.get("group_id")
    start = _parse_date(d.get("start_date"))
    end = _parse_date(d.get("end_date"))
    if not group_id or not start or not end:
        return jsonify({"error": "group_id, start_date, and end_date required"}), 400

    result = get_roster_calendar(group_id, start, end)
    rows = result.get("rows", [])
    group = RosterGroup.query.get(group_id)
    group_name = group.name if group else "roster"

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.formatting.rule import CellIsRule
    except ImportError:
        return jsonify({"error": "openpyxl not installed on server"}), 500

    wb = openpyxl.Workbook()
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _fill(hex_color):
        return PatternFill("solid", fgColor=hex_color.lstrip("#"))

    def _write_header(ws, headers, fill_color):
        hdr_font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
        for ci, h in enumerate(headers, 1):
            c = ws.cell(1, ci, h)
            c.font = hdr_font
            c.fill = _fill(fill_color)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border
        ws.row_dimensions[1].height = 25
        ws.freeze_panes = "A2"

    def _set_widths(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ── Sheet 1: Roster Schedule ──────────────────────────────
    ws_sched = wb.active
    ws_sched.title = "Roster Schedule"
    sched_headers = ["Date", "Day", "Employee", "Role", "Shift Start", "Shift End", "Status"]
    _write_header(ws_sched, sched_headers, "1F4973")

    alt_fill = _fill("EBF1FC")
    white_fill = _fill("FFFFFF")
    draft_font = Font(name="Arial", size=10, italic=True, color="B45309")
    normal_font = Font(name="Arial", size=10)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    for ri, r in enumerate(rows, 2):
        bg = alt_fill if ri % 2 == 0 else white_fill
        # Parse date for day-of-week
        try:
            dt = datetime.date.fromisoformat(r["date"])
            day_name = dt.strftime("%a")
        except (ValueError, AttributeError):
            day_name = ""

        vals = [r["date"], day_name, r["employee"], r["type"],
                r["start"], r["end"], r["status"]]
        is_draft = r.get("status", "").lower() == "draft"
        for ci, v in enumerate(vals, 1):
            c = ws_sched.cell(ri, ci, v or "")
            c.fill = bg
            c.border = border
            c.font = draft_font if is_draft else normal_font
            c.alignment = center if ci != 3 else left
        ws_sched.row_dimensions[ri].height = 20

    _set_widths(ws_sched, [12, 6, 20, 10, 12, 12, 12])

    # ── Sheet 2: Summary ──────────────────────────────────────
    ws_sum = wb.create_sheet("Summary")

    # Build per-employee stats from rows
    day_names = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    emp_stats = {}
    for r in rows:
        name = r.get("employee", "Unknown")
        if name not in emp_stats:
            emp_stats[name] = {
                "total": 0, "days": [0]*7, "sat": 0, "sun": 0,
                "draft": 0, "published": 0, "closings": 0,
                "openings": 0,
                "close_days": [0]*7,   # closings per weekday
                "open_days": [0]*7,    # openings per weekday
                "roles": {},
            }
        s = emp_stats[name]
        s["total"] += 1
        dow = None
        try:
            dt = datetime.date.fromisoformat(r["date"])
            dow = dt.weekday()  # 0=Mon
            s["days"][dow] += 1
            if dow == 5:
                s["sat"] += 1
            elif dow == 6:
                s["sun"] += 1
        except (ValueError, AttributeError):
            pass
        status = r.get("status", "").lower()
        if status == "draft":
            s["draft"] += 1
        else:
            s["published"] += 1
        role = r.get("type", "")
        s["roles"][role] = s["roles"].get(role, 0) + 1
        if role.lower() in ("closer", "close"):
            s["closings"] += 1
            if dow is not None:
                s["close_days"][dow] += 1
        if role.lower() in ("opener", "open"):
            s["openings"] += 1
            if dow is not None:
                s["open_days"][dow] += 1

    # Collect all unique roles for columns
    all_roles = sorted({role for s in emp_stats.values() for role in s["roles"]})
    sorted_names = sorted(emp_stats.keys())
    num_employees = len(emp_stats)
    grand_total = sum(s["total"] for s in emp_stats.values()) or 1
    pct_fmt = '0.0%'
    bold_font = Font(name="Arial", size=10, bold=True)
    sum_alt = _fill("EBF4E6")

    # ── Helper: write a mini-table with totals/avg rows ──
    def _write_section(ws, headers, data_fn, start_row, header_color,
                       pct_col=None, widths=None, cond_fmt_cols=None,
                       section_title=None):
        """Write header + data rows + TOTAL + AVG for a section.
        data_fn(name, stats) → list of values (no name col — prepended here).
        pct_col = 1-based column index for the % column (if any).
        Returns the next available row after this section.
        """
        # Section title row
        if section_title:
            tc = ws.cell(start_row, 1, section_title)
            tc.font = Font(name="Arial", size=12, bold=True, color=header_color)
            tc.alignment = left
            ws.row_dimensions[start_row].height = 28
            start_row += 1

        # Column header
        for ci, h in enumerate(headers, 1):
            c = ws.cell(start_row, ci, h)
            c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            c.fill = _fill(header_color)
            c.border = border
            c.alignment = center if ci > 1 else left
        ws.row_dimensions[start_row].height = 22

        # Data rows
        for ri_offset, name in enumerate(sorted_names):
            ri = start_row + 1 + ri_offset
            bg = sum_alt if ri_offset % 2 == 0 else white_fill
            vals = [name] + data_fn(name, emp_stats[name])
            for ci, v in enumerate(vals, 1):
                c = ws.cell(ri, ci, v)
                c.fill = bg
                c.border = border
                c.font = bold_font if ci == 1 else normal_font
                c.alignment = left if ci == 1 else center
                if pct_col and ci == pct_col:
                    c.number_format = pct_fmt
            ws.row_dimensions[ri].height = 20

        # Totals row
        t_row = start_row + 1 + num_employees
        ws.cell(t_row, 1, "TOTAL").font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        ws.cell(t_row, 1).fill = _fill(header_color)
        ws.cell(t_row, 1).border = border
        ws.cell(t_row, 1).alignment = left
        data_start = start_row + 1
        for ci in range(2, len(headers) + 1):
            c = ws.cell(t_row, ci)
            if pct_col and ci == pct_col:
                c.value = 1.0
                c.number_format = pct_fmt
            else:
                c.value = sum(ws.cell(r, ci).value or 0 for r in range(data_start, t_row))
            c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            c.fill = _fill(header_color)
            c.border = border
            c.alignment = center
        ws.row_dimensions[t_row].height = 25

        # Averages row
        a_row = t_row + 1
        ws.cell(a_row, 1, "AVG / Employee").font = Font(name="Arial", size=10, bold=True, italic=True, color=header_color)
        ws.cell(a_row, 1).border = border
        ws.cell(a_row, 1).alignment = left
        ws.cell(a_row, 1).fill = white_fill
        for ci in range(2, len(headers) + 1):
            c = ws.cell(a_row, ci)
            if pct_col and ci == pct_col:
                c.value = 1.0 / num_employees if num_employees else 0
                c.number_format = pct_fmt
            else:
                total_val = ws.cell(t_row, ci).value or 0
                c.value = round(total_val / num_employees, 1) if num_employees else 0
            c.font = Font(name="Arial", size=10, italic=True, color=header_color)
            c.fill = white_fill
            c.border = border
            c.alignment = center
        ws.row_dimensions[a_row].height = 20

        # Conditional formatting — highlight high/low
        if cond_fmt_cols and num_employees > 1:
            low_cfill = PatternFill("solid", fgColor="C6EFCE")
            low_cfont = Font(name="Arial", size=10, color="006100")
            high_cfill = PatternFill("solid", fgColor="FFC7CE")
            high_cfont = Font(name="Arial", size=10, color="9C0006")
            last_data = start_row + num_employees
            for col_idx in cond_fmt_cols:
                col_letter = chr(ord("A") + col_idx - 1)
                rng = f"{col_letter}{data_start}:{col_letter}{last_data}"
                col_vals = [ws.cell(r, col_idx).value or 0 for r in range(data_start, last_data + 1)]
                if not col_vals:
                    continue
                avg_val = sum(col_vals) / len(col_vals)
                high_t = avg_val * 1.3 if avg_val > 0 else 999
                low_t = avg_val * 0.7 if avg_val > 0 else -1
                ws.conditional_formatting.add(
                    rng, CellIsRule(operator="greaterThan", formula=[str(round(high_t))],
                                    fill=high_cfill, font=high_cfont))
                ws.conditional_formatting.add(
                    rng, CellIsRule(operator="lessThan", formula=[str(round(low_t))],
                                    fill=low_cfill, font=low_cfont))

        if widths:
            _set_widths(ws, widths)

        return a_row + 2  # next section starts 2 rows after avg

    # ────────────────────────────────────────────────────────────
    # SECTION 1: Overall Shift Totals (green)
    # ────────────────────────────────────────────────────────────
    ov_headers = ["Employee"] + day_names + ["Total", "% of Total",
                  "Saturday", "Sunday", "Wknd Total", "Draft", "Published"]

    def _overall_data(name, s):
        pct = s["total"] / grand_total
        return (s["days"] + [s["total"], pct, s["sat"], s["sun"],
                s["sat"] + s["sun"], s["draft"], s["published"]])

    next_row = _write_section(
        ws_sum, ov_headers, _overall_data, 1, "548235",
        pct_col=10, cond_fmt_cols=list(range(2, 9)),
        widths=[20] + [6]*7 + [8, 10, 10, 10, 12, 8, 10],
        section_title="Overall Shift Totals")

    # ────────────────────────────────────────────────────────────
    # SECTION 2: Closings Breakdown (dark blue)
    # ────────────────────────────────────────────────────────────
    cl_headers = ["Employee"] + day_names + ["Total"]

    def _close_data(name, s):
        return s["close_days"] + [s["closings"]]

    next_row = _write_section(
        ws_sum, cl_headers, _close_data, next_row, "1F4E79",
        cond_fmt_cols=list(range(2, 9)),
        widths=[20] + [6]*7 + [8],
        section_title="Closings Breakdown")

    # ────────────────────────────────────────────────────────────
    # SECTION 3: Openings Breakdown (teal)
    # ────────────────────────────────────────────────────────────
    op_headers = ["Employee"] + day_names + ["Total"]

    def _open_data(name, s):
        return s["open_days"] + [s["openings"]]

    next_row = _write_section(
        ws_sum, op_headers, _open_data, next_row, "0D7377",
        cond_fmt_cols=list(range(2, 9)),
        widths=[20] + [6]*7 + [8],
        section_title="Openings Breakdown")

    # ────────────────────────────────────────────────────────────
    # SECTION 4: Weekends Breakdown (purple)
    # ────────────────────────────────────────────────────────────
    wk_headers = ["Employee", "Saturday", "Sunday", "Wknd Total"]

    def _wknd_data(name, s):
        return [s["sat"], s["sun"], s["sat"] + s["sun"]]

    next_row = _write_section(
        ws_sum, wk_headers, _wknd_data, next_row, "6B4C9A",
        cond_fmt_cols=[2, 3, 4],
        widths=[20, 10, 10, 12],
        section_title="Weekends Breakdown")

    # ────────────────────────────────────────────────────────────
    # SECTION 5: Role Totals (olive)
    # ────────────────────────────────────────────────────────────
    if all_roles:
        role_headers = ["Employee"] + all_roles + ["Total"]

        def _role_data(name, s):
            return [s["roles"].get(r, 0) for r in all_roles] + [s["total"]]

        _write_section(
            ws_sum, role_headers, _role_data, next_row, "7D6B3A",
            widths=[20] + [max(10, len(r) + 4) for r in all_roles] + [8],
            section_title="Role Totals")

    # ── Sheet 3: Stat Day Summary ─────────────────────────────
    # Find all special days (holidays / half-days) in the date range
    from sqlalchemy import and_
    special_days = SpecialDay.query.filter(
        SpecialDay.date >= start,
        SpecialDay.date <= end,
        SpecialDay.day_type.in_(["holiday", "half_day"]),
    ).order_by(SpecialDay.date).all()

    # Filter to ones that apply to this planning unit
    pu_id = group.planning_unit_id if group else None
    applicable_days = []
    for sd in special_days:
        if sd.applies_to_all:
            applicable_days.append(sd)
        elif pu_id and any(pu.id == pu_id for pu in sd.planning_units):
            applicable_days.append(sd)

    if applicable_days:
        ws_stat = wb.create_sheet("Stat Day Summary")
        # Build a lookup: date -> list of employee names working that day
        from collections import defaultdict
        stat_workers = defaultdict(list)
        for r in rows:
            try:
                dt = datetime.date.fromisoformat(r["date"])
            except (ValueError, AttributeError):
                continue
            stat_workers[dt].append({
                "employee": r.get("employee", "Unknown"),
                "role": r.get("type", ""),
                "start": r.get("start", ""),
                "end": r.get("end", ""),
            })

        # Header
        stat_headers = ["Stat Day", "Date", "Day", "Type", "Employee", "Role",
                        "Shift Start", "Shift End"]
        _write_header(ws_stat, stat_headers, "B91C1C")

        stat_alt = _fill("FEF2F2")
        ri = 2
        for sd in applicable_days:
            workers = stat_workers.get(sd.date, [])
            if workers:
                for w in workers:
                    bg = stat_alt if ri % 2 == 0 else white_fill
                    vals = [sd.name, sd.date.isoformat(),
                            sd.date.strftime("%a"), sd.day_type.title(),
                            w["employee"], w["role"], w["start"], w["end"]]
                    for ci, v in enumerate(vals, 1):
                        c = ws_stat.cell(ri, ci, v or "")
                        c.fill = bg
                        c.border = border
                        c.font = normal_font
                        c.alignment = center if ci != 5 else left
                    ws_stat.row_dimensions[ri].height = 20
                    ri += 1
            else:
                # Show the stat day with "No one scheduled"
                bg = stat_alt if ri % 2 == 0 else white_fill
                vals = [sd.name, sd.date.isoformat(),
                        sd.date.strftime("%a"), sd.day_type.title(),
                        "— No one scheduled —", "", "", ""]
                for ci, v in enumerate(vals, 1):
                    c = ws_stat.cell(ri, ci, v or "")
                    c.fill = bg
                    c.border = border
                    c.font = Font(name="Arial", size=10, italic=True, color="999999") if ci == 5 else normal_font
                    c.alignment = center if ci != 5 else left
                ws_stat.row_dimensions[ri].height = 20
                ri += 1

        # Add a per-employee stat day count summary below
        ri += 1
        tc = ws_stat.cell(ri, 1, "Employee Stat Day Counts")
        tc.font = Font(name="Arial", size=12, bold=True, color="B91C1C")
        tc.alignment = left
        ws_stat.row_dimensions[ri].height = 28
        ri += 1

        count_headers = ["Employee"] + [sd.name for sd in applicable_days] + ["Total"]
        for ci, h in enumerate(count_headers, 1):
            c = ws_stat.cell(ri, ci, h)
            c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            c.fill = _fill("B91C1C")
            c.border = border
            c.alignment = center if ci > 1 else left
        ws_stat.row_dimensions[ri].height = 22
        ri += 1

        # Count per employee per stat day
        emp_stat_counts = {}
        for sd in applicable_days:
            for w in stat_workers.get(sd.date, []):
                name = w["employee"]
                if name not in emp_stat_counts:
                    emp_stat_counts[name] = {}
                emp_stat_counts[name][sd.name] = emp_stat_counts[name].get(sd.name, 0) + 1

        for emp_name in sorted(emp_stat_counts.keys()):
            bg = stat_alt if ri % 2 == 0 else white_fill
            counts = emp_stat_counts[emp_name]
            vals = [emp_name] + [counts.get(sd.name, 0) for sd in applicable_days]
            total = sum(v for v in vals[1:])
            vals.append(total)
            for ci, v in enumerate(vals, 1):
                c = ws_stat.cell(ri, ci, v)
                c.fill = bg
                c.border = border
                c.font = bold_font if ci == 1 else normal_font
                c.alignment = left if ci == 1 else center
            ws_stat.row_dimensions[ri].height = 20
            ri += 1

        _set_widths(ws_stat, [20, 12, 6, 10, 20, 10, 12, 12])

    # ── Sheet 4: Info ─────────────────────────────────────────
    ws_info = wb.create_sheet("Info")
    _write_header(ws_info, ["Field", "Value"], "6B4C9A")
    info_rows = [
        ["Group", group_name],
        ["Period", f"{start} to {end}"],
        ["Generated", datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")],
        ["Total Assignments", len(rows)],
        ["Employees", len(emp_stats)],
        ["Powered by", "Serevo — Bergeron WFM Solutions"],
    ]

    # Check for shift time anomalies (start_time > end_time, likely overnight)
    from app.models import RosterSlot
    anomaly_slots = []
    group_slots = RosterSlot.query.filter_by(roster_group_id=group_id).all()
    for sl in group_slots:
        if sl.shift_start and sl.shift_end and sl.shift_start > sl.shift_end:
            anomaly_slots.append(f"{sl.label or sl.role} ({sl.day_type}): {sl.shift_start.strftime('%H:%M')}→{sl.shift_end.strftime('%H:%M')}")
    if anomaly_slots:
        info_rows.append(["", ""])
        info_rows.append(["⚠ Shift Time Warnings", "Slots where start > end (possible overnight or config issue)"])
        for a in anomaly_slots:
            info_rows.append(["", a])

    info_alt = _fill("F5F0FF")
    warn_fill = _fill("FFF3CD")
    for ri, (field, val) in enumerate(info_rows, 2):
        is_warning = field.startswith("⚠") if field else (ri > len(info_rows) - len(anomaly_slots) + 1 and anomaly_slots)
        bg = warn_fill if is_warning else (info_alt if ri % 2 == 0 else white_fill)
        for ci, v in enumerate([field, val], 1):
            c = ws_info.cell(ri, ci, v)
            c.fill = bg
            c.border = border
            c.font = bold_font if ci == 1 else normal_font
            c.alignment = left
    _set_widths(ws_info, [25, 55])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_name = group_name.replace(" ", "_")
    filename = f"roster_{safe_name}_{start}_{end}.xlsx"
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ── Helpers ──────────────────────────────────────────────────

def _parse_time(v):
    if not v:
        return None
    if isinstance(v, datetime.time):
        return v
    try:
        parts = v.split(":")
        return datetime.time(int(parts[0]), int(parts[1]))
    except (ValueError, AttributeError, IndexError):
        return None


def _parse_date(v):
    if not v:
        return None
    if isinstance(v, datetime.date):
        return v
    try:
        return datetime.date.fromisoformat(v)
    except (ValueError, AttributeError):
        return None
