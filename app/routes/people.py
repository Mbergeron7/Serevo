"""
routes/people.py — People Management blueprint
================================================
Employee roster, accommodations, and PTO management.
"""

import json
import logging
import datetime as _dt
from datetime import datetime

from flask import (Blueprint, render_template, request, redirect,
                   url_for, jsonify, abort)
from app.auth import login_required, admin_required, get_current_user
from app.models import (db, EmployeeAvailability, Schedule, Contract,
                        Employee, EmployeePlanningUnit,
                        EmployeeContract, Selection, SelectionMember,
                        SkillMapping, PlanningUnit, ShiftSequence,
                        EmployeeShiftSequence, CoachingSession,
                        TimeClock, EmployeeWorkTimePattern, WorkTimePatternModel,
                        EmployeeQuartile, WeekTimePattern,
                        AttendanceEvent, AdherenceException,
                        PTOEntry, TimeOffType,
                        LeaveOfAbsence, LEAVE_TYPES, Accommodation)

log = logging.getLogger("serevo.people")

people_bp = Blueprint("people", __name__, url_prefix="/people")



from app.routes._utils import get_sheet as _get_sheet


# ── One-time sync: legacy planning_unit_id from junction table ──
def _sync_planning_units():
    """Sync Employee.planning_unit_id and schedule records to match
    current (active today) junction records.  Returns count of employees touched."""
    updated = 0
    _today = _dt.date.today()
    current_assignments = EmployeePlanningUnit.query.filter(
        EmployeePlanningUnit.valid_to.is_(None),
        db.or_(
            EmployeePlanningUnit.valid_from.is_(None),
            EmployeePlanningUnit.valid_from <= _today,
        ),
    ).all()
    for a in current_assignments:
        emp = Employee.query.get(a.employee_id)
        if not emp:
            continue
        effective = a.valid_from or _today
        touched = False
        # Sync legacy FK
        if emp.planning_unit_id != a.planning_unit_id:
            emp.planning_unit_id = a.planning_unit_id
            touched = True
        # Always fix schedule records from effective date onward
        n = Schedule.query.filter(
            Schedule.employee_id == emp.id,
            Schedule.schedule_date >= effective,
            Schedule.planning_unit_id != a.planning_unit_id,
        ).update({Schedule.planning_unit_id: a.planning_unit_id}, synchronize_session="fetch")
        if n:
            touched = True
        if touched:
            updated += 1
    return updated


@people_bp.route("/api/sync-planning-units", methods=["POST"])
@login_required
def sync_planning_units():
    updated = _sync_planning_units()
    db.session.commit()
    return jsonify({"success": True, "updated": updated})


# ── Roster View ──────────────────────────────────────────────
@people_bp.route("/")
@login_required
def roster():
    user = get_current_user()

    # Demo mode: serve fake data, no API/sheet calls
    if user and user.get("is_demo"):
        from app.demo_data import (get_demo_employees, get_demo_accommodations,
                                   get_demo_pto, _DEMO_ROTATION, DEMO_EMPLOYEES)
        employees, emp_err = get_demo_employees()
        accoms, _ = get_demo_accommodations()
        pto_list, _ = get_demo_pto()

        # Build demo avail_map — all demo employees have availability set
        avail_map = {}
        for e in employees:
            full = f"{e.get('First Name', '')} {e.get('Last Name', '')}".strip()
            if full:
                avail_map[full] = True

        # Build demo schedule_map — employees with schedules created
        # In demo, the 4 rotation employees have generated schedules
        _sched_ids = set(_DEMO_ROTATION.keys())
        schedule_map = {}
        for e in employees:
            if e.get("Employee ID") in _sched_ids:
                full = f"{e.get('First Name', '')} {e.get('Last Name', '')}".strip()
                schedule_map[full] = True
    else:
        from app.people.manager import get_employees, get_accommodations, get_pto
        sheet = _get_sheet()
        employees, emp_err = get_employees()  # always read from DB (edits go there)
        accoms, _ = get_accommodations(sheet)
        pto_list, _ = get_pto(sheet)

    # Build lookup maps for the template
    accom_map = {}
    for a in accoms:
        name = a.get("Employee", "").strip()
        if name:
            accom_map[name] = a

    pto_map = {}
    for p in pto_list:
        name = p.get("Employee", "").strip()
        if name:
            pto_map.setdefault(name, []).append(p)

    # Build availability & rotation maps from DB (non-demo only)
    if not (user and user.get("is_demo")):
        try:
            avail_emp_ids = set(
                r[0] for r in db.session.query(EmployeeAvailability.employee_id).distinct().all()
            )
        except Exception:
            db.session.rollback()
            avail_emp_ids = set()
        avail_map = {}
        for e in employees:
            db_id = e.get("_db_id")
            if db_id and db_id in avail_emp_ids:
                full = f"{e.get('First Name', '')} {e.get('Last Name', '')}".strip()
                avail_map[full] = True

        sched_emp_ids = set(
            r[0] for r in db.session.query(Schedule.employee_id).distinct().all()
        )
        schedule_map = {}
        for e in employees:
            db_id = e.get("_db_id")
            if db_id and db_id in sched_emp_ids:
                full = f"{e.get('First Name', '')} {e.get('Last Name', '')}".strip()
                schedule_map[full] = True

    # Count stats
    active = [e for e in employees
              if str(e.get("Status", "")).strip().lower() in ("active", "")]
    # LOB dropdown — single canonical source
    all_pu_names = PlanningUnit.get_available_lob_names()

    rotation_map = {}  # placeholder — rotation feature not yet wired

    if user and user.get("is_demo"):
        contracts = []
    else:
        try:
            contracts = [c.to_dict() for c in Contract.query.order_by(Contract.name).all()]
        except Exception:
            db.session.rollback()
            contracts = []

    return render_template("people/roster.html",
        user=user,
        employees=employees,
        accom_map=accom_map,
        pto_map=pto_map,
        avail_map=avail_map,
        schedule_map=schedule_map,
        emp_error=emp_err,
        active_count=len(active),
        total_count=len(employees),
        lob_count=len(lobs),
        accom_count=len(accoms),
        pto_count=len(pto_list),
        all_pu_names=all_pu_names,
        contracts=contracts,
        rotation_map=rotation_map,
    )


# ── Export ───────────────────────────────────────────────────
@people_bp.route("/export")
@login_required
def export_roster():
    """Download the employee roster as an Excel file."""
    import io
    from flask import send_file
    user = get_current_user()

    if user and user.get("is_demo"):
        from app.demo_data import get_demo_employees
        employees, _ = get_demo_employees()
    else:
        from app.people.manager import get_employees
        employees, _ = get_employees()

    # Build rows
    headers = ["Employee ID", "First Name", "Last Name", "Email",
               "Planning Unit", "Skill Group", "Team Lead", "Status",
               "Start Date", "Contract"]
    rows = []
    for e in employees:
        rows.append([
            e.get("Employee ID", ""),
            e.get("First Name", ""),
            e.get("Last Name", ""),
            e.get("Email", ""),
            e.get("Latest Skill Name", ""),
            e.get("Skill Group", ""),
            e.get("Team Lead", ""),
            e.get("Status", "Active"),
            e.get("Start Date", ""),
            e.get("Contract", ""),
        ])

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Employee Roster"

        # Header row
        header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # Data rows
        data_font = Font(name="Arial", size=10)
        for r, row in enumerate(rows, 2):
            for c, val in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=val)
                cell.font = data_font

        # Auto-width columns
        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 30)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                         as_attachment=True, download_name="employee_roster.xlsx")
    except Exception:
        log.exception("Export failed, falling back to CSV")
        import csv
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)
        buf = io.BytesIO(output.getvalue().encode("utf-8"))
        return send_file(buf, mimetype="text/csv", as_attachment=True,
                         download_name="employee_roster.csv")


# ── Accommodations ───────────────────────────────────────────
@people_bp.route("/accommodations")
@login_required
def accommodations():
    user = get_current_user()

    if user and user.get("is_demo"):
        from app.demo_data import get_demo_accommodations, get_demo_employees
        accoms, err = get_demo_accommodations()
        emps, _ = get_demo_employees()
        emp_names = [f"{e['First Name']} {e['Last Name']}" for e in emps]
    else:
        from app.people.manager import get_accommodations, get_employee_names
        sheet = _get_sheet()
        accoms, err = get_accommodations(sheet)
        emp_names, _ = get_employee_names(sheet)

    return render_template("people/accommodations.html",
        user=user,
        accommodations=accoms,
        employee_names=emp_names,
        error=err,
    )


# ── Leave of Absence ─────────────────────────────────────────
@people_bp.route("/loa")
@login_required
def loa():
    user = get_current_user()

    if user and user.get("is_demo"):
        from app.demo_data import get_demo_loa, get_demo_employees
        loa_records = get_demo_loa()
        demo_emps, _ = get_demo_employees()
        emp_list = [{"id": i + 1, "name": f'{e["First Name"]} {e["Last Name"]}'} for i, e in enumerate(demo_emps) if e.get("Status") == "Active"]
        return render_template("people/loa.html",
            user=user,
            loa_records=loa_records,
            leave_types=LEAVE_TYPES,
            emp_list=emp_list,
        )

    loa_records = []
    try:
        all_loa = LeaveOfAbsence.query.order_by(LeaveOfAbsence.start_date.desc()).all()
        for rec in all_loa:
            d = rec.to_dict()
            d["employee_name"] = rec.employee.full_name if rec.employee else "Unknown"
            loa_records.append(d)
    except Exception:
        db.session.rollback()

    emp_list = []
    try:
        emps = Employee.query.filter(Employee.status != "Terminated").order_by(Employee.last_name).all()
        emp_list = [{"id": e.id, "name": e.full_name} for e in emps]
    except Exception:
        pass

    return render_template("people/loa.html",
        user=user,
        loa_records=loa_records,
        leave_types=LEAVE_TYPES,
        emp_list=emp_list,
    )


@people_bp.route("/accommodations/save", methods=["POST"])
@login_required
def save_accommodation():
    from app.people.manager import save_accommodation as _save
    try:
        payload = request.get_json(silent=True) or {}
        err = _save(payload)
        if err:
            return jsonify({"success": False, "error": err})
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})


@people_bp.route("/accommodations/delete", methods=["POST"])
@login_required
def delete_accommodation():
    from app.people.manager import delete_accommodation as _delete
    try:
        payload = request.get_json(silent=True) or {}
        name = payload.get("employee", "")
        err = _delete(name)
        if err:
            return jsonify({"success": False, "error": err})
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})


# ── PTO ──────────────────────────────────────────────────────
@people_bp.route("/pto")
@login_required
def pto():
    user = get_current_user()
    return render_template("people/pto.html", user=user)


@people_bp.route("/pto/save", methods=["POST"])
@login_required
def save_pto():
    from app.people.manager import save_pto as _save
    try:
        payload = request.get_json(silent=True) or {}
        err = _save(payload)
        if err:
            return jsonify({"success": False, "error": err})
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})


@people_bp.route("/pto/delete", methods=["POST"])
@login_required
def delete_pto():
    from app.people.manager import delete_pto as _delete
    try:
        payload = request.get_json(silent=True) or {}
        row_idx = int(payload.get("row_index", 0))
        if row_idx < 2:
            return jsonify({"success": False, "error": "Invalid row"})
        err = _delete(row_idx)
        if err:
            return jsonify({"success": False, "error": err})
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})


# ── Availability ────────────────────────────────────────────
# ── API: employee names (for other modules to use) ───────────
@people_bp.route("/api/names")
@login_required
def api_names():
    # Return DB employees with IDs for dropdowns
    emps = Employee.query.filter(Employee.status.in_(["Active", "active"])).order_by(Employee.last_name, Employee.first_name).all()
    if emps:
        return jsonify({"names": [{"id": e.id, "name": e.full_name} for e in emps]})
    # Fallback to sheets-based names
    from app.people.manager import get_employee_names
    names, err = get_employee_names()
    if err:
        return jsonify({"names": [], "error": err})
    return jsonify({"names": names})


def _demo_guard():
    """Return a mock-success JSON response if the current user is a demo user, else None."""
    user = get_current_user()
    if user and user.get("is_demo"):
        return jsonify(ok=True, demo=True, message="Changes are not saved in demo mode.")


# ═══════════════════════════════════════════════════════════════
# EMPLOYEE PROFILE PAGE
# ═══════════════════════════════════════════════════════════════

@people_bp.route("/<emp_id>/profile")
@login_required
def employee_profile(emp_id):
    user = get_current_user()

    # Demo mode: build a mock employee object from DEMO_EMPLOYEES
    if user and user.get("is_demo"):
        from app.demo_data import DEMO_EMPLOYEES
        demo_emp = None
        for e in DEMO_EMPLOYEES:
            if str(e.get("Employee ID", "")) == str(emp_id):
                demo_emp = e
                break
        if not demo_emp:
            abort(404)

        class _MockPU:
            def __init__(self, name): self.name = name
        class _MockEmp:
            def __init__(self, d):
                self.id = d.get("Employee ID")
                self.employee_id = d.get("Employee ID", "")
                self.first_name = d.get("First Name", "")
                self.last_name = d.get("Last Name", "")
                self.full_name = f"{self.first_name} {self.last_name}".strip()
                self.status = d.get("Status", "Active")
                self.team_lead = d.get("Team Lead", "")
                self.languages = d.get("Languages", "English")
                from app.utils.timezone import get_app_timezone_name
                self.timezone = d.get("Timezone", get_app_timezone_name())
                self.skill_start = d.get("Latest Skill Start", "")
                self.skill_end = d.get("Latest Skill End", "")
                self.end_date = d.get("End Date", "")
                self.schedule_excluded = d.get("Schedule Excluded", "") == "Yes"
                self.email = d.get("Email", "")
                self.external_id_1 = ""
                self.external_id_2 = ""
                lob = d.get("Latest Skill Name", "")
                self.planning_unit = _MockPU(lob) if lob else None
                self.contract = None

        emp = _MockEmp(demo_emp)
        return render_template("people/profile.html",
            user=user,
            emp=emp,
            skill_mappings=[],
            skill_history=[],
            contract_assignments=[],
            contract_history=[],
            selection_memberships=[],
            selection_history=[],
            pu_assignments=[],
            pu_history=[],
            ss_assignments=[],
            ss_history=[],
            wtp_assignments=[],
            quartile=None,
            quartile_history=[],
            quartile_entities=[],
            avail_map={},
        )

    # emp_id may be a numeric DB id or a string employee_id like "E1001"
    try:
        numeric_id = int(emp_id)
        emp = Employee.query.get(numeric_id)
    except (ValueError, TypeError):
        emp = None

    # Fallback: look up by employee_id string
    if emp is None:
        emp = Employee.query.filter_by(employee_id=str(emp_id)).first()

    if emp is None:
        abort(404)

    # Skills — split active (no valid_to) vs history (has valid_to)
    skill_mappings = SkillMapping.query.filter_by(
        employee_id=emp.id
    ).filter(SkillMapping.valid_to.is_(None)).order_by(
        SkillMapping.valid_from.desc().nullslast()).all()
    skill_history = SkillMapping.query.filter_by(
        employee_id=emp.id
    ).filter(SkillMapping.valid_to.isnot(None)).order_by(
        SkillMapping.valid_to.desc()).all()

    # Contracts — split active vs history
    contract_assignments = EmployeeContract.query.filter_by(
        employee_id=emp.id
    ).filter(EmployeeContract.valid_to.is_(None)).all()
    contract_history = EmployeeContract.query.filter_by(
        employee_id=emp.id
    ).filter(EmployeeContract.valid_to.isnot(None)).order_by(
        EmployeeContract.valid_to.desc()).all()

    # Selections — split active vs history
    selection_memberships = SelectionMember.query.filter_by(
        employee_id=emp.id
    ).filter(SelectionMember.valid_to.is_(None)).all()
    selection_history = SelectionMember.query.filter_by(
        employee_id=emp.id
    ).filter(SelectionMember.valid_to.isnot(None)).order_by(
        SelectionMember.valid_to.desc()).all()

    # Planning Units — split active vs history
    pu_assignments = EmployeePlanningUnit.query.filter_by(
        employee_id=emp.id
    ).filter(EmployeePlanningUnit.valid_to.is_(None)).order_by(
        EmployeePlanningUnit.priority).all()
    pu_history = EmployeePlanningUnit.query.filter_by(
        employee_id=emp.id
    ).filter(EmployeePlanningUnit.valid_to.isnot(None)).order_by(
        EmployeePlanningUnit.valid_to.desc()).all()
    try:
        ss_assignments = EmployeeShiftSequence.query.filter_by(
            employee_id=emp.id
        ).filter(EmployeeShiftSequence.valid_to.is_(None)).all()
        ss_history = EmployeeShiftSequence.query.filter_by(
            employee_id=emp.id
        ).filter(EmployeeShiftSequence.valid_to.isnot(None)).order_by(
            EmployeeShiftSequence.valid_to.desc()).all()
    except Exception:
        db.session.rollback()
        ss_assignments = []
        ss_history = []
    try:
        wtp_assignments = EmployeeWorkTimePattern.query.filter_by(employee_id=emp.id).all()
    except Exception:
        db.session.rollback()
        wtp_assignments = []
    try:
        quartile = EmployeeQuartile.query.filter_by(
            employee_id=emp.id
        ).filter(EmployeeQuartile.end_date.is_(None)).first()
        quartile_history = EmployeeQuartile.query.filter_by(
            employee_id=emp.id
        ).filter(EmployeeQuartile.end_date.isnot(None)).order_by(
            EmployeeQuartile.end_date.desc()).all()
    except Exception:
        db.session.rollback()
        quartile = None
        quartile_history = []
    # Build availability map keyed by day_of_week (0-6)
    avail_entries = EmployeeAvailability.query.filter_by(employee_id=emp.id).all()
    avail_map = {a.day_of_week: a for a in avail_entries}

    # Load quartile entities for the dropdown
    try:
        quartile_entities = WeekTimePattern.query.filter_by(is_active=True).order_by(WeekTimePattern.name).all()
    except Exception:
        quartile_entities = []

    return render_template("people/profile.html",
        user=user,
        emp=emp,
        skill_mappings=skill_mappings,
        skill_history=skill_history,
        contract_assignments=contract_assignments,
        contract_history=contract_history,
        selection_memberships=selection_memberships,
        selection_history=selection_history,
        pu_assignments=pu_assignments,
        pu_history=pu_history,
        ss_assignments=ss_assignments,
        ss_history=ss_history,
        wtp_assignments=wtp_assignments,
        quartile=quartile,
        quartile_history=quartile_history,
        quartile_entities=quartile_entities,
        avail_map=avail_map,
    )


# ═══════════════════════════════════════════════════════════════
# PROFILE API ENDPOINTS
# ═══════════════════════════════════════════════════════════════

@people_bp.route("/api/profile/contracts-list")
@login_required
def profile_contracts_list():
    items = Contract.query.filter_by(is_active=True).order_by(Contract.name).all()
    return jsonify({"items": [{"id": c.id, "name": c.name} for c in items]})


@people_bp.route("/api/profile/selections-list")
@login_required
def profile_selections_list():
    items = Selection.query.filter_by(is_active=True).order_by(Selection.name).all()
    return jsonify({"items": [{"id": s.id, "name": s.name} for s in items]})


@people_bp.route("/api/profile/shift-sequences-list")
@login_required
def profile_shift_sequences_list():
    items = ShiftSequence.query.filter_by(is_active=True).order_by(ShiftSequence.name).all()
    return jsonify({"items": [{"id": s.id, "name": s.name} for s in items]})


@people_bp.route("/api/profile/rotation-models-list")
@login_required
def profile_rotation_models_list():
    items = WorkTimePatternModel.query.filter_by(is_active=True).order_by(WorkTimePatternModel.name).all()
    return jsonify({"items": [{"id": m.id, "name": m.name} for m in items]})


@people_bp.route("/api/profile/assign", methods=["POST"])
@login_required
def profile_assign():
    dg = _demo_guard()
    if dg:
        return dg
    data = request.get_json(silent=True) or {}
    assign_type = data.get("type")
    emp_id = data.get("employee_id")
    item_id = data.get("item_id")
    assignment_id = data.get("assignment_id")  # if editing existing

    if not emp_id or (not item_id and not assignment_id and not data.get("schedule_quartile_id")):
        return jsonify({"success": False, "error": "Missing employee or item ID."})

    valid_from = _dt.datetime.strptime(data["valid_from"], "%Y-%m-%d").date() if data.get("valid_from") else None
    valid_to = _dt.datetime.strptime(data["valid_to"], "%Y-%m-%d").date() if data.get("valid_to") else None

    try:
        _today = _dt.date.today()
        effective = valid_from or _today  # when the new assignment takes effect

        if assign_type == "planning_unit":

            if assignment_id:
                obj = EmployeePlanningUnit.query.get(int(assignment_id))
                if not obj:
                    return jsonify({"success": False, "error": "Assignment not found."})
                old_pu_id_before = obj.planning_unit_id
                obj.planning_unit_id = item_id
                obj.priority = int(data.get("priority", 1))
                obj.valid_from = valid_from
                obj.valid_to = valid_to
                # Keep legacy direct FK in sync if this assignment is currently active
                if not valid_to and effective <= _today:
                    emp = Employee.query.get(emp_id)
                    if emp and emp.planning_unit_id != item_id:
                        emp.planning_unit_id = item_id
                # Transfer schedules from effective date onward
                if old_pu_id_before != item_id:
                    Schedule.query.filter(
                        Schedule.employee_id == emp_id,
                        Schedule.schedule_date >= effective,
                    ).update({Schedule.planning_unit_id: item_id}, synchronize_session="fetch")
            else:
                # End-date any current assignments — day before new one starts
                end_date = effective - _dt.timedelta(days=1)
                current = EmployeePlanningUnit.query.filter_by(
                    employee_id=emp_id
                ).filter(EmployeePlanningUnit.valid_to.is_(None)).all()
                for c in current:
                    c.valid_to = end_date
                obj = EmployeePlanningUnit(
                    employee_id=emp_id, planning_unit_id=item_id,
                    priority=int(data.get("priority", 1)),
                    valid_from=effective, valid_to=valid_to,
                )
                db.session.add(obj)
                # Keep legacy direct FK in sync if effective today or earlier
                if effective <= _today:
                    emp = Employee.query.get(emp_id)
                    if emp:
                        emp.planning_unit_id = item_id
                # Transfer schedules from effective date onward
                Schedule.query.filter(
                    Schedule.employee_id == emp_id,
                    Schedule.schedule_date >= effective,
                ).update({Schedule.planning_unit_id: item_id}, synchronize_session="fetch")
        elif assign_type == "contract":
            if assignment_id:
                obj = EmployeeContract.query.get(int(assignment_id))
                if not obj:
                    return jsonify({"success": False, "error": "Assignment not found."})
                obj.contract_id = item_id
                obj.valid_from = valid_from
                obj.valid_to = valid_to
            else:
                # End-date any current assignments — day before new one starts
                current = EmployeeContract.query.filter_by(
                    employee_id=emp_id
                ).filter(EmployeeContract.valid_to.is_(None)).all()
                for c in current:
                    c.valid_to = effective - _dt.timedelta(days=1)
                obj = EmployeeContract(
                    employee_id=emp_id, contract_id=item_id,
                    valid_from=effective, valid_to=valid_to,
                )
                db.session.add(obj)
        elif assign_type == "selection":
            if assignment_id:
                obj = SelectionMember.query.get(int(assignment_id))
                if not obj:
                    return jsonify({"success": False, "error": "Assignment not found."})
                obj.selection_id = item_id
                obj.valid_from = valid_from
                obj.valid_to = valid_to
            else:
                # End-date any current membership — day before new one starts
                current = SelectionMember.query.filter_by(
                    employee_id=emp_id, selection_id=item_id
                ).filter(SelectionMember.valid_to.is_(None)).all()
                for c in current:
                    c.valid_to = effective - _dt.timedelta(days=1)
                obj = SelectionMember(
                    selection_id=item_id, employee_id=emp_id,
                    valid_from=effective, valid_to=valid_to,
                )
                db.session.add(obj)
        elif assign_type == "shift_sequence":
            ref_date = _dt.datetime.strptime(data["reference_date"], "%Y-%m-%d").date() if data.get("reference_date") else None
            row_index = int(data.get("row_index", 0))
            if assignment_id:
                obj = EmployeeShiftSequence.query.get(int(assignment_id))
                if not obj:
                    return jsonify({"success": False, "error": "Assignment not found."})
                obj.shift_sequence_id = item_id
                obj.row_index = row_index
                obj.reference_date = ref_date
                obj.valid_from = valid_from
                obj.valid_to = valid_to
            else:
                # End-date any current assignments — day before new one starts
                current = EmployeeShiftSequence.query.filter_by(
                    employee_id=emp_id
                ).filter(EmployeeShiftSequence.valid_to.is_(None)).all()
                for c in current:
                    c.valid_to = effective - _dt.timedelta(days=1)
                obj = EmployeeShiftSequence(
                    employee_id=emp_id, shift_sequence_id=item_id,
                    row_index=row_index,
                    reference_date=ref_date,
                    valid_from=effective, valid_to=valid_to,
                )
                db.session.add(obj)
        elif assign_type == "work_time_pattern":
            ref_date = _dt.datetime.strptime(data["reference_date"], "%Y-%m-%d").date() if data.get("reference_date") else None
            if assignment_id:
                obj = EmployeeWorkTimePattern.query.get(int(assignment_id))
                if not obj:
                    return jsonify({"success": False, "error": "Assignment not found."})
                obj.work_time_pattern_model_id = item_id
                obj.reference_date = ref_date
                obj.valid_from = valid_from
                obj.valid_to = valid_to
            else:
                # End-date any current assignments — day before new one starts
                current = EmployeeWorkTimePattern.query.filter_by(
                    employee_id=emp_id
                ).filter(EmployeeWorkTimePattern.valid_to.is_(None)).all()
                for c in current:
                    c.valid_to = effective - _dt.timedelta(days=1)
                obj = EmployeeWorkTimePattern(
                    employee_id=emp_id, work_time_pattern_model_id=item_id,
                    reference_date=ref_date,
                    valid_from=effective, valid_to=valid_to,
                )
                db.session.add(obj)
        elif assign_type == "quartile":
            sq_id = data.get("schedule_quartile_id")
            sq_id = int(sq_id) if sq_id else None
            # Derive planning_unit_id from the quartile entity
            derived_pu = None
            if sq_id:
                wtp = WeekTimePattern.query.get(sq_id)
                if wtp:
                    derived_pu = wtp.planning_unit_id
            if assignment_id:
                # Editing an existing quartile record
                obj = EmployeeQuartile.query.get(int(assignment_id))
                if not obj:
                    return jsonify({"success": False, "error": "Assignment not found."})
                obj.schedule_quartile_id = sq_id
                obj.planning_unit_id = derived_pu
                obj.quartile = int(data.get("quartile", 0)) or 0
                eff = data.get("effective_date")
                if eff:
                    obj.effective_date = _dt.datetime.strptime(eff, "%Y-%m-%d").date()
                obj.notes = data.get("notes", "")
            else:
                # End-date any current quartile — day before new one starts
                eff = data.get("effective_date")
                q_effective = _dt.datetime.strptime(eff, "%Y-%m-%d").date() if eff else _today
                current = EmployeeQuartile.query.filter_by(
                    employee_id=emp_id
                ).filter(EmployeeQuartile.end_date.is_(None)).all()
                for c in current:
                    c.end_date = q_effective - _dt.timedelta(days=1)
                obj = EmployeeQuartile(
                    employee_id=emp_id,
                    schedule_quartile_id=sq_id,
                    planning_unit_id=derived_pu,
                    quartile=int(data.get("quartile", 0)) or 0,
                    effective_date=q_effective,
                    notes=data.get("notes", ""),
                )
                db.session.add(obj)
        elif assign_type == "skill":
            if assignment_id:
                obj = SkillMapping.query.get(int(assignment_id))
                if not obj:
                    return jsonify({"success": False, "error": "Assignment not found."})
                if item_id:
                    obj.skill_group_id = item_id
                obj.proficiency = int(data.get("proficiency", 3))
                obj.priority = int(data.get("priority", 1))
                obj.valid_from = valid_from
                obj.valid_to = valid_to
            else:
                # End-date any current mapping — day before new one starts
                current = SkillMapping.query.filter_by(
                    employee_id=emp_id, skill_group_id=item_id
                ).filter(SkillMapping.valid_to.is_(None)).all()
                for c in current:
                    c.valid_to = effective - _dt.timedelta(days=1)
                obj = SkillMapping(
                    employee_id=emp_id, skill_group_id=item_id,
                    proficiency=int(data.get("proficiency", 3)),
                    priority=int(data.get("priority", 1)),
                    is_active=True,
                    valid_from=effective, valid_to=valid_to,
                )
                db.session.add(obj)
        else:
            return jsonify({"success": False, "error": f"Unknown assignment type: {assign_type}"})

        db.session.commit()
        return jsonify({"success": True})
    except Exception as e:
        db.session.rollback()
        err = str(e)
        if "uq_emp_wtpm" in err or "uq_sel_emp" in err or "uq_skill_employee" in err:
            return jsonify({"success": False, "error": "This assignment already exists."})
        return jsonify({"success": False, "error": err})


@people_bp.route("/api/profile/unassign", methods=["POST"])
@login_required
def profile_unassign():
    dg = _demo_guard()
    if dg:
        return dg
    data = request.get_json(silent=True) or {}
    assign_type = data.get("type")
    assign_id = data.get("assignment_id")

    if not assign_id:
        return jsonify({"success": False, "error": "Missing assignment ID."})

    model_map = {
        "planning_unit": EmployeePlanningUnit,
        "contract": EmployeeContract,
        "selection": SelectionMember,
        "shift_sequence": EmployeeShiftSequence,
        "skill": SkillMapping,
        "work_time_pattern": EmployeeWorkTimePattern,
        "quartile": EmployeeQuartile,
    }
    model = model_map.get(assign_type)
    if not model:
        return jsonify({"success": False, "error": f"Unknown type: {assign_type}"})

    obj = model.query.get(assign_id)
    if not obj:
        return jsonify({"success": False, "error": "Assignment not found."})

    # For history-tracked types, end-date instead of deleting
    _today = _dt.date.today()
    if assign_type in ("planning_unit", "shift_sequence", "skill", "selection") and hasattr(obj, "valid_to"):
        obj.valid_to = _today
        # If end-dating a planning unit, update legacy FK to next active or clear it
        if assign_type == "planning_unit":
            emp = Employee.query.get(obj.employee_id)
            if emp:
                # Find another current (non-end-dated) assignment for this employee
                other = EmployeePlanningUnit.query.filter(
                    EmployeePlanningUnit.employee_id == obj.employee_id,
                    EmployeePlanningUnit.id != obj.id,
                    EmployeePlanningUnit.valid_to.is_(None),
                ).first()
                emp.planning_unit_id = other.planning_unit_id if other else None
    elif assign_type == "quartile" and hasattr(obj, "end_date"):
        obj.end_date = _today
    else:
        db.session.delete(obj)
    db.session.commit()
    return jsonify({"success": True})


@people_bp.route("/api/profile/delete-assignment", methods=["POST"])
@login_required
def profile_delete_assignment():
    dg = _demo_guard()
    if dg:
        return dg
    data = request.get_json(silent=True) or {}
    assign_type = data.get("type")
    assign_id = data.get("assignment_id")

    if not assign_id:
        return jsonify({"success": False, "error": "Missing assignment ID."})

    model_map = {
        "planning_unit": EmployeePlanningUnit,
        "contract": EmployeeContract,
        "selection": SelectionMember,
        "shift_sequence": EmployeeShiftSequence,
        "skill": SkillMapping,
        "work_time_pattern": EmployeeWorkTimePattern,
        "quartile": EmployeeQuartile,
    }
    model = model_map.get(assign_type)
    if not model:
        return jsonify({"success": False, "error": f"Unknown type: {assign_type}"})

    obj = model.query.get(assign_id)
    if not obj:
        return jsonify({"success": False, "error": "Assignment not found."})

    db.session.delete(obj)
    db.session.commit()
    return jsonify({"success": True})


@people_bp.route("/api/profile/availability", methods=["POST"])
@login_required
def profile_availability():
    dg = _demo_guard()
    if dg:
        return dg
    data = request.get_json(silent=True) or {}
    emp_id = data.get("employee_id")
    days = data.get("days", [])

    if not emp_id:
        return jsonify({"success": False, "error": "Missing employee ID."})

    try:
        for day in days:
            dow = day["day_of_week"]
            existing = EmployeeAvailability.query.filter_by(
                employee_id=emp_id, day_of_week=dow).first()

            earliest = _dt.datetime.strptime(day["earliest_start"], "%H:%M").time() if day.get("earliest_start") else None
            latest = _dt.datetime.strptime(day["latest_end"], "%H:%M").time() if day.get("latest_end") else None

            if existing:
                existing.is_available = day.get("is_available", True)
                existing.earliest_start = earliest
                existing.latest_end = latest
            else:
                obj = EmployeeAvailability(
                    employee_id=emp_id,
                    day_of_week=dow,
                    is_available=day.get("is_available", True),
                    earliest_start=earliest,
                    latest_end=latest,
                )
                db.session.add(obj)

        db.session.commit()
        return jsonify({"success": True})
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "error": str(e)})


# ── Coaching Sessions ─────────────────────────────────────────

@people_bp.route("/api/coaching/list", methods=["POST"])
@login_required
def coaching_list():
    data = request.get_json(silent=True) or {}
    emp_id = data.get("employee_id")
    user = get_current_user()

    if user and user.get("is_demo"):
        from app.demo_data import get_demo_coaching_sessions
        return jsonify({"sessions": get_demo_coaching_sessions(emp_id)})

    if not emp_id:
        return jsonify({"sessions": []})

    try:
        sessions = (CoachingSession.query
                    .filter_by(employee_id=emp_id)
                    .order_by(CoachingSession.session_date.desc())
                    .all())
        return jsonify({"sessions": [s.to_dict() for s in sessions]})
    except Exception as e:
        db.session.rollback()
        return jsonify({"sessions": [], "error": str(e)})


@people_bp.route("/api/coaching/save", methods=["POST"])
@login_required
@admin_required
def coaching_save():
    dg = _demo_guard()
    if dg:
        return dg
    data = request.get_json(silent=True) or {}
    session_id = data.get("id")

    try:
        if session_id:
            cs = CoachingSession.query.get(session_id)
            if not cs:
                return jsonify({"success": False, "error": "Session not found."})
        else:
            cs = CoachingSession(employee_id=data["employee_id"])
            db.session.add(cs)

        cs.coach_name = data.get("coach_name", "")
        cs.session_date = _dt.date.fromisoformat(data["session_date"])
        cs.session_time = data.get("session_time") or None
        cs.duration_mins = int(data.get("duration_mins") or 30)
        cs.topic = data.get("topic", "")
        cs.category = data.get("category", "general")
        cs.quality_score = float(data["quality_score"]) if data.get("quality_score") else None
        cs.notes = data.get("notes", "")
        cs.outcome = data.get("outcome", "")
        cs.follow_up = data.get("follow_up", "")

        db.session.commit()
        return jsonify({"success": True, "session": cs.to_dict()})
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "error": str(e)})


@people_bp.route("/api/coaching/delete", methods=["POST"])
@login_required
@admin_required
def coaching_delete():
    dg = _demo_guard()
    if dg:
        return dg
    data = request.get_json(silent=True) or {}
    session_id = data.get("id")

    cs = CoachingSession.query.get(session_id)
    if not cs:
        return jsonify({"success": False, "error": "Session not found."})

    db.session.delete(cs)
    db.session.commit()
    return jsonify({"success": True})


# ── Leave of Absence CRUD ─────────────────────────────────────────

@people_bp.route("/api/loa/list", methods=["POST"])
@login_required
def loa_list():
    data = request.get_json(silent=True) or {}
    emp_id = data.get("employee_id")
    if not emp_id:
        return jsonify({"success": False, "error": "Missing employee_id"})
    records = LeaveOfAbsence.query.filter_by(employee_id=int(emp_id)).order_by(
        LeaveOfAbsence.start_date.desc()).all()
    return jsonify({"success": True, "items": [r.to_dict() for r in records]})


@people_bp.route("/api/loa/save", methods=["POST"])
@login_required
@admin_required
def loa_save():
    dg = _demo_guard()
    if dg:
        return dg
    data = request.get_json(silent=True) or {}
    loa_id = data.get("id")

    if loa_id:
        rec = LeaveOfAbsence.query.get(int(loa_id))
        if not rec:
            return jsonify({"success": False, "error": "Record not found"})
    else:
        rec = LeaveOfAbsence(employee_id=int(data["employee_id"]))
        db.session.add(rec)

    rec.leave_type = data.get("leave_type", "Medical")
    sd = data.get("start_date")
    rec.start_date = _dt.date.fromisoformat(sd) if sd else _dt.date.today()
    erd = data.get("expected_return_date")
    rec.expected_return_date = _dt.date.fromisoformat(erd) if erd else None
    ard = data.get("actual_return_date")
    rec.actual_return_date = _dt.date.fromisoformat(ard) if ard else None
    rec.status = data.get("status", "Active")
    rec.notes = data.get("notes", "")

    db.session.commit()
    return jsonify({"success": True, "item": rec.to_dict()})


@people_bp.route("/api/loa/delete", methods=["POST"])
@login_required
@admin_required
def loa_delete():
    dg = _demo_guard()
    if dg:
        return dg
    data = request.get_json(silent=True) or {}
    loa_id = data.get("id")
    rec = LeaveOfAbsence.query.get(int(loa_id))
    if not rec:
        return jsonify({"success": False, "error": "Record not found"})
    db.session.delete(rec)
    db.session.commit()
    return jsonify({"success": True})


# ── Expiry Warnings (LOA + Accommodations) ────────────────────────

@people_bp.route("/api/expiry-warnings")
@login_required
def expiry_warnings():
    """Return LOA and accommodation records expiring within 14 days."""
    today = _dt.date.today()
    horizon = today + _dt.timedelta(days=14)
    warnings = []

    # LOA records with expected_return_date approaching
    try:
        loas = LeaveOfAbsence.query.filter(
            LeaveOfAbsence.status == "Active",
            LeaveOfAbsence.expected_return_date.isnot(None),
            LeaveOfAbsence.expected_return_date <= horizon,
            LeaveOfAbsence.expected_return_date >= today,
        ).all()
        for loa in loas:
            emp = Employee.query.get(loa.employee_id)
            days_left = (loa.expected_return_date - today).days
            warnings.append({
                "type": "loa",
                "employee": emp.full_name if emp else f"ID {loa.employee_id}",
                "employee_id": loa.employee_id,
                "detail": f"{loa.leave_type} leave — expected return {loa.expected_return_date.isoformat()}",
                "days_left": days_left,
                "record_id": loa.id,
            })
    except Exception:
        db.session.rollback()

    # Accommodation records with end_date approaching
    try:
        accoms = Accommodation.query.filter(
            Accommodation.end_date.isnot(None),
            Accommodation.end_date <= horizon,
            Accommodation.end_date >= today,
        ).all()
        for acc in accoms:
            emp = Employee.query.get(acc.employee_id)
            days_left = (acc.end_date - today).days
            warnings.append({
                "type": "accommodation",
                "employee": emp.full_name if emp else f"ID {acc.employee_id}",
                "employee_id": acc.employee_id,
                "detail": f"Accommodation ending {acc.end_date.isoformat()}",
                "days_left": days_left,
                "record_id": acc.id,
            })
    except Exception:
        db.session.rollback()

    warnings.sort(key=lambda w: w["days_left"])
    return jsonify({"success": True, "warnings": warnings})


# ── Attendance history for profile ───────────────────────────────

@people_bp.route("/api/profile/attendance", methods=["POST"])
@login_required
def profile_attendance():
    """Return recent attendance entries for an employee profile."""
    user = get_current_user()
    if not user or user.get("role") not in ("admin", "supervisor"):
        return jsonify(success=False, error="Forbidden"), 403

    if user.get("is_demo"):
        from app.demo_data import get_demo_profile_attendance
        payload = request.get_json(silent=True) or {}
        data = get_demo_profile_attendance(payload.get("employee_id"), int(payload.get("days", 30)))
        return jsonify(success=True, **data)

    payload = request.get_json(silent=True) or {}
    employee_id = payload.get("employee_id")
    days = int(payload.get("days", 30))

    if not employee_id:
        return jsonify(success=False, error="Missing employee_id")

    from datetime import date, timedelta
    from sqlalchemy import func

    cutoff = date.today() - timedelta(days=days)

    entries = TimeClock.query.filter(
        TimeClock.employee_id == employee_id,
        TimeClock.date >= cutoff,
    ).order_by(TimeClock.clock_in.desc()).limit(100).all()

    rows = []
    for e in entries:
        rows.append({
            "id": e.id,
            "date": e.date.isoformat() if e.date else "",
            "clock_in": e.clock_in.strftime("%H:%M") if e.clock_in else "",
            "clock_out": e.clock_out.strftime("%H:%M") if e.clock_out else "",
            "total_hours": e.total_hours,
            "status": e.status,
        })

    # Summary stats
    completed = [e for e in entries if e.status == "completed"]
    total_hours = sum(e.total_hours or 0 for e in completed)
    avg_hours = round(total_hours / len(completed), 1) if completed else 0

    # ── Attendance events (coded by supervisor) ──
    from datetime import date, timedelta as _td
    cutoff_date = date.today() - _td(days=days)
    events = AttendanceEvent.query.filter(
        AttendanceEvent.employee_id == employee_id,
        AttendanceEvent.event_date >= cutoff_date,
    ).order_by(AttendanceEvent.event_date.desc()).all()
    event_rows = [ev.to_dict() for ev in events]

    # Counts by category
    cat_counts = {}
    for ev in events:
        exc = ev.adherence_exception
        cat = exc.category if exc else "other"
        cat_counts[cat] = cat_counts.get(cat, 0) + 1

    # ── PTO entries ──
    pto_entries = PTOEntry.query.filter(
        PTOEntry.employee_id == employee_id,
        PTOEntry.start_date >= cutoff_date,
    ).order_by(PTOEntry.start_date.desc()).all()
    pto_rows = []
    for p in pto_entries:
        tot = p.time_off_type
        pto_rows.append({
            "id": p.id,
            "start_date": p.start_date.isoformat(),
            "end_date": p.end_date.isoformat(),
            "pto_type": p.pto_type,
            "time_off_type": tot.label if tot else "Time Off",
            "time_off_color": tot.color if tot else "#6b7280",
            "approval_status": p.approval_status,
            "note": p.note or "",
        })

    return jsonify(
        success=True,
        rows=rows,
        events=event_rows,
        pto=pto_rows,
        summary={
            "total_entries": len(entries),
            "completed": len(completed),
            "total_hours": round(total_hours, 1),
            "avg_hours_per_shift": avg_hours,
            "days": days,
            "absences": cat_counts.get("absence", 0),
            "lates": cat_counts.get("late", 0),
            "early_outs": cat_counts.get("early_out", 0),
            "other_events": cat_counts.get("other", 0) + cat_counts.get("break_overrun", 0),
            "pto_count": len([p for p in pto_entries if p.approval_status == "approved"]),
            "pto_pending": len([p for p in pto_entries if p.approval_status == "pending"]),
        },
    )


# ── Attendance Event CRUD ───────────────────────────────────────

@people_bp.route("/api/attendance-events", methods=["POST"])
@login_required
def list_attendance_events():
    """List attendance events, optionally filtered by employee."""
    user = get_current_user()
    if not user or user.get("role") not in ("admin", "supervisor"):
        return jsonify(success=False, error="Forbidden"), 403

    payload = request.get_json(silent=True) or {}
    employee_id = payload.get("employee_id")
    days = int(payload.get("days", 30))

    from datetime import date, timedelta
    cutoff = date.today() - timedelta(days=days)

    q = AttendanceEvent.query.filter(AttendanceEvent.event_date >= cutoff)
    if employee_id:
        q = q.filter(AttendanceEvent.employee_id == employee_id)
    events = q.order_by(AttendanceEvent.event_date.desc()).limit(200).all()

    return jsonify(success=True, events=[e.to_dict() for e in events])


@people_bp.route("/api/attendance-events/save", methods=["POST"])
@login_required
def save_attendance_event():
    """Create or update an attendance event."""
    user = get_current_user()
    if not user or user.get("role") not in ("admin", "supervisor"):
        return jsonify(success=False, error="Forbidden"), 403

    payload = request.get_json(silent=True) or {}
    event_id = payload.get("id")
    employee_id = payload.get("employee_id")
    event_date_str = payload.get("event_date")
    exception_id = payload.get("adherence_exception_id")

    if not employee_id or not event_date_str or not exception_id:
        return jsonify(success=False, error="Missing required fields")

    from datetime import date
    try:
        event_date = date.fromisoformat(event_date_str)
    except (ValueError, TypeError):
        return jsonify(success=False, error="Invalid date format")

    # Check the exception code exists
    exc = AdherenceException.query.get(exception_id)
    if not exc:
        return jsonify(success=False, error="Invalid adherence code")

    from app.models import User
    uid = None
    if user.get("username"):
        u = User.query.filter_by(username=user["username"]).first()
        uid = u.id if u else None

    if event_id:
        ev = AttendanceEvent.query.get(event_id)
        if not ev:
            return jsonify(success=False, error="Event not found")
        ev.employee_id = employee_id
        ev.event_date = event_date
        ev.adherence_exception_id = exception_id
        ev.note = payload.get("note", "")
        ev.is_excused = bool(payload.get("is_excused", exc.is_excused))
    else:
        ev = AttendanceEvent(
            employee_id=employee_id,
            event_date=event_date,
            adherence_exception_id=exception_id,
            note=payload.get("note", ""),
            is_excused=bool(payload.get("is_excused", exc.is_excused)),
            coded_by=uid,
        )
        db.session.add(ev)

    db.session.commit()
    return jsonify(success=True, event=ev.to_dict())


@people_bp.route("/api/attendance-events/delete", methods=["POST"])
@login_required
def delete_attendance_event():
    """Delete an attendance event."""
    user = get_current_user()
    if not user or user.get("role") not in ("admin", "supervisor"):
        return jsonify(success=False, error="Forbidden"), 403

    payload = request.get_json(silent=True) or {}
    event_id = payload.get("id")
    if not event_id:
        return jsonify(success=False, error="Missing event id")

    ev = AttendanceEvent.query.get(event_id)
    if not ev:
        return jsonify(success=False, error="Event not found")

    db.session.delete(ev)
    db.session.commit()
    return jsonify(success=True)


@people_bp.route("/api/adherence-codes", methods=["POST"])
@login_required
def list_adherence_codes():
    """Return active adherence exception codes for dropdowns."""
    codes = AdherenceException.query.filter_by(is_active=True)\
        .order_by(AdherenceException.sort_order, AdherenceException.label).all()
    return jsonify(success=True, codes=[c.to_dict() for c in codes])


# ── PTO / Time Off (DB-backed) ──────────────────────────────────

@people_bp.route("/api/pto-entries", methods=["POST"])
@login_required
def list_pto_entries():
    """Return PTO entries, optionally filtered."""
    user = get_current_user()
    if not user or user.get("role") not in ("admin", "supervisor"):
        return jsonify(success=False, error="Forbidden"), 403

    payload = request.get_json(silent=True) or {}
    employee_id = payload.get("employee_id")
    status_filter = payload.get("status")  # pending | approved | denied
    days = int(payload.get("days", 90))

    from datetime import date, timedelta
    cutoff = date.today() - timedelta(days=days)

    q = PTOEntry.query.filter(PTOEntry.start_date >= cutoff)
    if employee_id:
        q = q.filter(PTOEntry.employee_id == employee_id)
    if status_filter:
        q = q.filter(PTOEntry.approval_status == status_filter)

    entries = q.order_by(PTOEntry.start_date.desc()).limit(200).all()

    rows = []
    for p in entries:
        emp = Employee.query.get(p.employee_id)
        tot = p.time_off_type
        pu = emp.planning_unit if emp else None
        rows.append({
            "id": p.id,
            "employee_id": p.employee_id,
            "employee_name": emp.full_name if emp else "",
            "group": pu.name if pu else "",
            "start_date": p.start_date.isoformat(),
            "end_date": p.end_date.isoformat(),
            "pto_type": p.pto_type,
            "time_off_type": tot.label if tot else "Time Off",
            "time_off_color": tot.color if tot else "#6b7280",
            "approval_status": p.approval_status,
            "note": p.note or "",
            "updated_at": p.updated_at.strftime("%Y-%m-%d %H:%M") if p.updated_at else "",
        })

    return jsonify(success=True, entries=rows)


@people_bp.route("/api/pto-entries/save", methods=["POST"])
@login_required
def save_pto_entry():
    """Create a PTO entry (admin/supervisor adding directly)."""
    user = get_current_user()
    if not user or user.get("role") not in ("admin", "supervisor"):
        return jsonify(success=False, error="Forbidden"), 403

    payload = request.get_json(silent=True) or {}
    employee_id = payload.get("employee_id")
    start_str = payload.get("start_date")
    end_str = payload.get("end_date")
    time_off_type_id = payload.get("time_off_type_id")

    if not employee_id or not start_str or not end_str:
        return jsonify(success=False, error="Missing required fields")

    from datetime import date
    try:
        start = date.fromisoformat(start_str)
        end = date.fromisoformat(end_str)
    except (ValueError, TypeError):
        return jsonify(success=False, error="Invalid date format")

    from app.models import User
    uid = None
    if user.get("username"):
        u = User.query.filter_by(username=user["username"]).first()
        uid = u.id if u else None

    entry = PTOEntry(
        employee_id=employee_id,
        start_date=start,
        end_date=end,
        pto_type=payload.get("pto_type", "full"),
        time_off_type_id=time_off_type_id,
        approval_status="approved",  # admin-added = auto-approved
        requested_by=uid,
        reviewed_by=uid,
        note=payload.get("note", ""),
    )
    db.session.add(entry)
    db.session.commit()
    return jsonify(success=True)


@people_bp.route("/api/pto-entries/delete", methods=["POST"])
@login_required
def delete_pto_entry():
    """Delete a PTO entry."""
    user = get_current_user()
    if not user or user.get("role") not in ("admin", "supervisor"):
        return jsonify(success=False, error="Forbidden"), 403

    payload = request.get_json(silent=True) or {}
    entry_id = payload.get("id")
    if not entry_id:
        return jsonify(success=False, error="Missing entry id")

    entry = PTOEntry.query.get(entry_id)
    if not entry:
        return jsonify(success=False, error="Entry not found")

    db.session.delete(entry)
    db.session.commit()
    return jsonify(success=True)


@people_bp.route("/api/time-off-types", methods=["POST"])
@login_required
def list_time_off_types():
    """Return active time-off types for dropdowns."""
    types = TimeOffType.query.filter_by(is_active=True)\
        .order_by(TimeOffType.sort_order, TimeOffType.label).all()
    return jsonify(success=True, types=[t.to_dict() for t in types])
